"""Syntetická „nepořádná“ složka zakázky pro cvičení Kancelář (kancelar.qmd).

Spuštění (z kořene workshop/):
    uv run --with reportlab --with pypdf --with openpyxl --with python-docx --with pillow --with matplotlib python tools/make_office_data.py

Vytvoří:
    data/kancelar/zakazka_neporadek.zip   ~55 malých souborů tří zakázek promíchaných dohromady
    data/kancelar/seznam_zakazek.csv      seznam zakázek kanceláře (čísla, názvy, přezdívky)
    demos/kancelar/klic_zarazeni.csv      správné zařazení každého souboru (klíč k řešení)

Vše je smyšlené (firmy, osoby, adresy). Semínko 2026, výstup je bajtově reprodukovatelný
(PDF přes reportlab invariant=1, pevná data v ZIP, pevné metadata Office souborů).
Písmo DejaVuSans se bere z balíčku matplotlib (kvůli české diakritice v PDF a fotkách).
"""
from __future__ import annotations

import csv
import io
import random
import re
import zipfile
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import matplotlib
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "kancelar"
DEMO = ROOT / "demos" / "kancelar"
TOP = "zakazka_neporadek/"
rng = random.Random(2026)
CET = timezone(timedelta(hours=1))  # pevná zóna → reprodukovatelné e-maily

FONT_FILE = Path(matplotlib.get_data_path()) / "fonts/ttf/DejaVuSans.ttf"
FONT_BOLD = Path(matplotlib.get_data_path()) / "fonts/ttf/DejaVuSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("DV", str(FONT_FILE)))
pdfmetrics.registerFont(TTFont("DVB", str(FONT_BOLD)))

KANCELAR = "STATIKA & DETAIL s.r.o. (fiktivní kancelář), Brno"
ZAKAZKY = {
    "Z2026-014": dict(nazev="Hala Modřice", slozka="Z2026-014_Hala_Modrice",
                      klient="Stavby Novák s.r.o. (fiktivní)", kontakt="Ing. Petra Malá",
                      misto="Modřice, průmyslová zóna", typ="ocelová hala 36 × 18 m",
                      stav="rozpracováno", zodpovedny="Ing. Karel Svoboda",
                      prezdivky="Modřice; Modrice; hala Modřice; HM"),
    "Z2026-009": dict(nazev="Lávka přes Svitavu", slozka="Z2026-009_Lavka_pres_Svitavu",
                      klient="Město Svitávka (fiktivní)", kontakt="Bc. Jana Horáková",
                      misto="Svitávka", typ="ocelová lávka pro pěší, rozpětí 24 m",
                      stav="rozpracováno", zodpovedny="Ing. Karel Svoboda",
                      prezdivky="lávka; lavka; Svitava; Svitavu"),
    "Z2025-031": dict(nazev="Parkovací dům Bystrc", slozka="Z2025-031_Parkovaci_dum_Bystrc",
                      klient="Parking Bystrc a.s. (fiktivní)", kontakt="Ing. Ondřej Beneš",
                      misto="Brno-Bystrc", typ="prefabrikovaný železobeton",
                      stav="dokončeno, fakturace", zodpovedny="Ing. Lucie Kopecká",
                      prezdivky="PD Bystrc; Bystrc; parkovací dům"),
}

# ---------------------------------------------------------------- helpers
FIXED = (2026, 3, 20, 8, 0, 0)


def normalize_zip(data: bytes) -> bytes:
    """Re-pack an Office file (xlsx/docx are ZIPs) with fixed timestamps → reproducible bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, FIXED[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps "now" into modified
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-03-01T08:00:00Z", data)
            dst.writestr(zi, data)
    return out.getvalue()


PAPER = {"A1": (841, 594), "A2": (594, 420), "A3": (420, 297), "A4": (297, 210)}


def axis_bubble(c, x, y, label):
    c.circle(x, y, 4 * mm)
    c.setFont("DV", 8)
    c.drawCentredString(x, y - 1.2 * mm, label)


def sketch(c, kind, x0, y0, w, h):
    """Very simple schematic 'drawing' so the PDF looks like a výkres."""
    c.setLineWidth(0.6)
    if kind == "kotveni":
        nx, ny = 7, 3
        for i in range(nx):
            x = x0 + i * w / (nx - 1)
            c.setDash(4, 2); c.line(x, y0 - 8 * mm, x, y0 + h + 4 * mm); c.setDash()
            axis_bubble(c, x, y0 - 12 * mm, str(i + 1))
            for j in range(ny):
                y = y0 + j * h / (ny - 1)
                c.rect(x - 6 * mm, y - 6 * mm, 12 * mm, 12 * mm)
                for dx in (-3, 3):
                    for dy in (-3, 3):
                        c.circle(x + dx * mm, y + dy * mm, 0.8 * mm)
        for j in range(ny):
            y = y0 + j * h / (ny - 1)
            c.setDash(4, 2); c.line(x0 - 8 * mm, y, x0 + w + 4 * mm, y); c.setDash()
            axis_bubble(c, x0 - 12 * mm, y, "ABC"[j])
    elif kind in ("sloupy", "oplasteni"):
        n = 7
        for i in range(n):
            x = x0 + i * w / (n - 1)
            c.setLineWidth(1.4); c.line(x, y0, x, y0 + h); c.setLineWidth(0.6)
            axis_bubble(c, x, y0 - 10 * mm, str(i + 1))
        c.setLineWidth(1.2); c.line(x0, y0 + h, x0 + w, y0 + h); c.setLineWidth(0.6)
        if kind == "sloupy":
            for i in (0, 4):
                xa, xb = x0 + i * w / 6, x0 + (i + 1) * w / 6
                c.line(xa, y0, xb, y0 + h); c.line(xa, y0 + h, xb, y0)
        else:
            for k in range(1, 6):
                c.line(x0, y0 + k * h / 6, x0 + w, y0 + k * h / 6)
    elif kind in ("strecha", "vaznice"):
        if kind == "strecha":
            c.setLineWidth(1.2)
            c.line(x0, y0, x0 + w, y0)
            c.line(x0, y0, x0 + w / 2, y0 + h * 0.6); c.line(x0 + w / 2, y0 + h * 0.6, x0 + w, y0)
            c.setLineWidth(0.6)
            for k in range(1, 12):
                xk = x0 + k * w / 12
                top = y0 + h * 0.6 * (1 - abs(xk - (x0 + w / 2)) / (w / 2))
                c.line(xk, y0, xk, top)
                if k < 11:
                    c.line(xk, y0, x0 + (k + 1) * w / 12, y0 + h * 0.6 * (1 - abs(x0 + (k + 1) * w / 12 - (x0 + w / 2)) / (w / 2)))
        else:
            c.rect(x0, y0, w, h)
            for k in range(1, 9):
                c.line(x0, y0 + k * h / 9, x0 + w, y0 + k * h / 9)
            for i in range(0, 6, 2):
                xa, xb = x0 + i * w / 6, x0 + (i + 1) * w / 6
                c.line(xa, y0, xb, y0 + h); c.line(xa, y0 + h, xb, y0)
    elif kind == "detaily":
        for k in range(3):
            cx = x0 + (k + 0.5) * w / 3
            c.rect(cx - 30 * mm, y0 + 10 * mm, 60 * mm, h - 20 * mm)
            for r in range(4):
                for s in range(2):
                    c.circle(cx - 10 * mm + s * 20 * mm, y0 + 25 * mm + r * (h - 50 * mm) / 3, 2 * mm)
            c.setFont("DV", 9)
            c.drawCentredString(cx, y0 + 2 * mm, f"Detail {'ABCDEF'[k + rng.randrange(2) * 3]}")
    elif kind == "schodiste":
        steps = 14
        for k in range(steps):
            xa = x0 + k * w / steps
            ya = y0 + k * h / steps
            c.line(xa, ya, xa, ya + h / steps)
            c.line(xa, ya + h / steps, xa + w / steps, ya + h / steps)
        c.setLineWidth(1.2); c.line(x0, y0 - 6 * mm, x0 + w, y0 + h - 6 * mm)
    elif kind == "lavka":
        c.setLineWidth(1.2)
        c.line(x0, y0 + h * 0.3, x0 + w, y0 + h * 0.3)
        c.line(x0, y0 + h * 0.75, x0 + w, y0 + h * 0.75)
        c.setLineWidth(0.6)
        for k in range(13):
            xa = x0 + k * w / 12
            c.line(xa, y0 + h * 0.3, xa, y0 + h * 0.75)
            if k < 12:
                c.line(xa, y0 + h * 0.3, xa + w / 12, y0 + h * 0.75)
        c.line(x0 - 10 * mm, y0, x0 + 15 * mm, y0); c.line(x0 + w - 15 * mm, y0, x0 + w + 10 * mm, y0)
    elif kind == "prefa":
        for i in range(5):
            for j in range(3):
                c.rect(x0 + i * w / 5 + 3 * mm, y0 + j * h / 3 + 3 * mm, w / 5 - 6 * mm, h / 3 - 6 * mm)
                c.setFont("DV", 7)
                c.drawString(x0 + i * w / 5 + 6 * mm, y0 + j * h / 3 + 6 * mm, f"P{j + 1}{i + 1:02d}")


def drawing_pdf(zak, cislo, nazev, rev, revize_hist, kind, fmt, meritko, datum, pages=1):
    W, H = PAPER[fmt]
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(W * mm, H * mm), invariant=1)
    c.setTitle(f"{cislo} {nazev}")
    c.setAuthor(KANCELAR)
    z = ZAKAZKY[zak]
    for p in range(pages):
        c.setLineWidth(1.5)
        c.rect(10 * mm, 10 * mm, (W - 20) * mm, (H - 20) * mm)
        # title block (bottom right) 180 x 62 mm
        bx, by, bw, bh = (W - 10 - 180) * mm, 10 * mm, 180 * mm, 62 * mm
        c.setLineWidth(0.8)
        c.rect(bx, by, bw, bh)
        rows = [
            ("Projektant", KANCELAR),
            ("Investor", z["klient"]),
            ("Akce", f"{z['nazev']}, {z['misto']}"),
            ("Zakázka", zak),
            ("Výkres", nazev),
            ("Číslo výkresu", cislo + (f"  (list {p + 1}/{pages})" if pages > 1 else "")),
            ("Revize", rev),
            ("Datum", datum),
            ("Měřítko / formát", f"{meritko} / {fmt}"),
        ]
        rh = bh / len(rows)
        for i, (k, v) in enumerate(rows):
            y = by + bh - (i + 1) * rh
            c.line(bx, y, bx + bw, y)
            c.setFont("DV", 6.5); c.drawString(bx + 2 * mm, y + 2 * mm, k.upper())
            c.setFont("DVB" if k in ("Číslo výkresu", "Revize", "Zakázka") else "DV", 8.5)
            c.drawString(bx + 34 * mm, y + 2 * mm, v)
        c.line(bx + 32 * mm, by, bx + 32 * mm, by + bh)
        # revision table above the title block
        ty = by + bh + 2 * mm
        c.setFont("DVB", 7); c.drawString(bx, ty + (len(revize_hist) + 0.4) * 5 * mm, "REVIZE  |  DATUM  |  POPIS ZMĚNY")
        c.setFont("DV", 7)
        for i, (r, d, popis) in enumerate(reversed(revize_hist)):
            c.drawString(bx, ty + i * 5 * mm, f"{r}   {d}   {popis}")
        # the picture
        sketch(c, kind, 40 * mm, 45 * mm, (W - 260) * mm, (H - 90) * mm)
        c.setFont("DVB", 14)
        c.drawString(20 * mm, (H - 22) * mm, f"{cislo} – {nazev.upper()}")
        c.showPage()
    c.save()
    return buf.getvalue()


def text_pdf(title, lines, author=KANCELAR, pages=1):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(210 * mm, 297 * mm), invariant=1)
    c.setTitle(title)
    c.setAuthor(author)
    per = (len(lines) + pages - 1) // pages
    for p in range(pages):
        c.setFont("DVB", 15); c.drawString(20 * mm, 270 * mm, title)
        c.setFont("DV", 10)
        y = 255 * mm
        for line in lines[p * per:(p + 1) * per]:
            c.drawString(20 * mm, y, line)
            y -= 6 * mm
        c.setFont("DV", 8); c.drawRightString(190 * mm, 12 * mm, f"strana {p + 1}/{pages}")
        c.showPage()
    c.save()
    return buf.getvalue()


def xlsx(rows_by_sheet, title):
    wb = Workbook()
    wb.properties.creator = "kancelář"
    wb.properties.title = title
    wb.properties.created = wb.properties.modified = datetime(2026, 3, 1, 8, 0, 0)
    first = True
    for sheet, rows in rows_by_sheet.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = sheet
        first = False
        for r in rows:
            ws.append(r)
        ws["A1"].font = Font(bold=True, size=13)
        for col in "ABCDEF":
            ws.column_dimensions[col].width = 18
    b = io.BytesIO()
    wb.save(b)
    return normalize_zip(b.getvalue())


def eml(frm, to, subject, date, body, cc=None):
    m = EmailMessage()
    m["From"] = frm
    m["To"] = to
    if cc:
        m["Cc"] = cc
    m["Subject"] = subject
    m["Date"] = format_datetime(date)
    m["Message-ID"] = f"<{date:%Y%m%d%H%M}.{rng.randrange(10**6):06d}@mail.example>"
    m.set_content(body)
    return bytes(m)


def photo(caption, date, desc, seed):
    r = random.Random(seed)
    img = Image.new("RGB", (480, 360))
    d = ImageDraw.Draw(img)
    for y in range(200):  # sky
        g = 170 + y // 4
        d.line([(0, y), (480, y)], fill=(120 + y // 5, g, 230))
    d.rectangle([0, 200, 480, 360], fill=(120 + r.randrange(20), 110, 95))
    for _ in range(r.randrange(3, 7)):  # steel columns / precast
        x = r.randrange(20, 440)
        d.rectangle([x, 70 + r.randrange(60), x + 12, 260], fill=(70, 80, 95))
    d.rectangle([0, 230, 480, 250], fill=(150, 150, 150))
    for _ in range(400):  # noise so it looks a bit like a photo
        x, y = r.randrange(480), r.randrange(200, 360)
        d.point((x, y), fill=(r.randrange(80, 160),) * 3)
    font = ImageFont.truetype(str(FONT_FILE), 15)
    d.rectangle([250, 270, 470, 345], fill=(245, 245, 235), outline=(30, 30, 30), width=2)
    for i, line in enumerate(caption.split("\n")):
        d.text((258, 277 + i * 20), line, fill=(20, 20, 20), font=font)
    exif = Image.Exif()
    exif[0x010E] = desc  # ImageDescription
    exif[0x0132] = date.strftime("%Y:%m:%d %H:%M:%S")  # DateTime
    exif[0x0110] = "Telefon (syntetický)"  # Model
    b = io.BytesIO()
    img.save(b, "JPEG", quality=55, exif=exif.tobytes())
    return b.getvalue()


def docx_contract():
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 2, 10, 9, 0, 0)
    doc.core_properties.author = "kancelář"
    doc.add_heading("SMLOUVA O DÍLO – NÁVRH", 0)
    z = ZAKAZKY["Z2026-014"]
    for p in [
        f"Objednatel: {z['klient']}, zastoupený {z['kontakt']}",
        f"Zhotovitel: {KANCELAR}",
        f"Předmět díla: dílenská dokumentace ocelové konstrukce – akce {z['nazev']} (zakázka Z2026-014).",
        "Rozsah: 3D model Tekla Structures, výrobní a montážní výkresy, výkaz materiálu.",
        "Cena díla: 186 000 Kč bez DPH. Vícepráce sazbou 950 Kč/h bez DPH.",
        "Termín: odevzdání 1. vydání výkresů do 15. 3. 2026.",
        "NÁVRH K PŘIPOMÍNKÁM – není podepsáno.",
    ]:
        doc.add_paragraph(p)
    b = io.BytesIO()
    doc.save(b)
    return normalize_zip(b.getvalue())


def ifc(name, zak, date, n=12):
    lines = [
        "ISO-10303-21;", "HEADER;",
        "FILE_DESCRIPTION(('ViewDefinition [CoordinationView_V2.0]'),'2;1');",
        f"FILE_NAME('{name}','{date}',('detailer'),('STATIKA a DETAIL - fiktivni'),'','Tekla Structures (synteticky export)','');",
        "FILE_SCHEMA(('IFC2X3'));", "ENDSEC;", "DATA;",
        f"#1=IFCPROJECT('3vB2YO$MX4xv5uCqZZG05x',$,'{zak}','{ZAKAZKY[zak]['slozka'].split('_', 1)[1].replace('_', ' ')}',$,$,$,$,$);",
    ]
    for i in range(n):
        lines.append(f"#{i + 10}=IFCBEAM('{rng.getrandbits(64):016X}{i:06d}',$,'{'COLUMN' if i % 3 == 0 else 'BEAM'}','HEA{200 + 20 * (i % 4)}',$,$,$,'P{i + 1}');")
    lines += ["ENDSEC;", "END-ISO-10303-21;", ""]
    return "\n".join(lines).encode("ascii")


def dwg(note):
    return (b"AC1032" + bytes(122) + "\n".join([
        "ZASTUPNY SOUBOR PRO WORKSHOP - NENI TO SKUTECNY VYKRES DWG.",
        note, ""]).encode("ascii") + bytes(rng.getrandbits(8) for _ in range(512)))


# ---------------------------------------------------------------- content
FILES: list[tuple[str, bytes, tuple]] = []   # (path inside zip, bytes, mtime)
KEY: list[dict] = []


def add(path, data, when, zak="", typ="", cislo="", rev="", stav="aktualni", pozn=""):
    FILES.append((path, data, when))
    KEY.append(dict(puvodni_cesta=path, zakazka=zak, typ=typ, cislo=cislo, revize=rev, stav=stav, poznamka=pozn))


def T(y, m, d, h=9, mi=0):
    return (y, m, d, h, mi, 0)


Z = "Z2026-014"
HIST = {
    "V-101": [("A", "16.02.2026", "První vydání"), ("B", "02.03.2026", "Posun kotev v ose 4 dle zaměření")],
    "V-102": [("A", "18.02.2026", "První vydání"), ("B", "27.02.2026", "Změna profilu sloupů HEA 260 → HEA 280"),
              ("C", "09.03.2026", "Doplnění ztužidel v poli 5–6")],
    "V-103": [("A", "19.02.2026", "První vydání"), ("B", "28.02.2026", "Úprava sklonu střechy 5°"),
              ("C", "10.03.2026", "Doplnění světlíku")],
    "V-104": [("A", "11.03.2026", "První vydání")],
    "V-105": [("A", "20.02.2026", "První vydání"), ("B", "05.03.2026", "Změna rozteče paždíků")],
    "V-106": [("A", "06.03.2026", "První vydání")],
    "V-107": [("A", "08.03.2026", "První vydání")],
}
DRW = {  # cislo: (nazev, kind, format, meritko)
    "V-101": ("Kotevní plán", "kotveni", "A2", "1:100"),
    "V-102": ("Sloupy a ztužidla – osa A", "sloupy", "A1", "1:50"),
    "V-103": ("Střešní vazník", "strecha", "A1", "1:50"),
    "V-104": ("Vaznice a ztužení střechy", "vaznice", "A2", "1:100"),
    "V-105": ("Opláštění – paždíky", "oplasteni", "A2", "1:100"),
    "V-106": ("Detaily přípojů", "detaily", "A3", "1:10"),
    "V-107": ("Schodiště", "schodiste", "A3", "1:25"),
}


def drw(cislo, rev, zak=Z):
    nazev, kind, fmt, mer = DRW[cislo]
    hist = [h for h in HIST[cislo] if h[0] <= rev]
    return drawing_pdf(zak, cislo, nazev, rev, hist, kind, fmt, mer, hist[-1][1], pages=2 if cislo == "V-106" else 1)


def build():
    d = {}
    # --- Hala Modřice: drawings -------------------------------------------------
    d["V-101A"] = drw("V-101", "A"); d["V-101B"] = drw("V-101", "B")
    d["V-102A"] = drw("V-102", "A"); d["V-102C"] = drw("V-102", "C")
    d["V-103B"] = drw("V-103", "B"); d["V-103C"] = drw("V-103", "C")
    d["V-104A"] = drw("V-104", "A")
    d["V-105A"] = drw("V-105", "A"); d["V-105B"] = drw("V-105", "B")
    d["V-106A"] = drw("V-106", "A"); d["V-107A"] = drw("V-107", "A")

    add(TOP + "V-101 kotevni plan.pdf", d["V-101A"], T(2026, 2, 16, 15, 40), Z, "VYK", "V-101", "A", "stara_revize")
    add(TOP + "od Pavla/kotevni plan_opraveny.pdf", d["V-101B"], T(2026, 3, 2, 17, 5), Z, "VYK", "V-101", "B")
    add(TOP + "V-102_sloupy_revA.pdf", d["V-102A"], T(2026, 2, 18, 11, 20), Z, "VYK", "V-102", "A", "stara_revize")
    add(TOP + "V-102_sloupy_revC.pdf", d["V-102C"], T(2026, 3, 9, 16, 30), Z, "VYK", "V-102", "C")
    add(TOP + "V102 sloupy FINAL.pdf", d["V-102C"], T(2026, 3, 9, 16, 45), Z, "VYK", "V-102", "C", "duplicita",
        "stejný obsah jako V-102_sloupy_revC.pdf")
    add(TOP + "Nová složka/kopie - V-102_sloupy_revC.pdf", d["V-102C"], T(2026, 3, 11, 8, 10), Z, "VYK", "V-102", "C",
        "duplicita", "stejný obsah jako V-102_sloupy_revC.pdf")
    add(TOP + "vykres strecha final.pdf", d["V-103B"], T(2026, 2, 28, 18, 2), Z, "VYK", "V-103", "B", "stara_revize")
    add(TOP + "vykres strecha final_final2.pdf", d["V-103C"], T(2026, 3, 10, 19, 12), Z, "VYK", "V-103", "C")
    add(TOP + "Stažené soubory/V-103_strecha_revB.pdf", d["V-103C"], T(2026, 3, 12, 7, 55), Z, "VYK", "V-103", "C",
        "duplicita", "POZOR: název říká rev B, v rohovém razítku je rev C; stejný obsah jako final_final2")
    add(TOP + "Scan_20260312_0001.pdf", d["V-104A"], T(2026, 3, 12, 10, 3), Z, "VYK", "V-104", "A",
        pozn="název neříká nic, obsah poznáte z razítka")
    add(TOP + "V-105_B.pdf", d["V-105B"], T(2026, 3, 5, 14, 0), Z, "VYK", "V-105", "B")
    add(TOP + "stare/V-105_A.pdf", d["V-105A"], T(2026, 2, 20, 13, 30), Z, "VYK", "V-105", "A", "stara_revize")
    add(TOP + "detaily pripoju.pdf", d["V-106A"], T(2026, 3, 6, 12, 0), Z, "VYK", "V-106", "A", pozn="2 listy")
    add(TOP + "HM_schodiste_V107_A.pdf", d["V-107A"], T(2026, 3, 8, 9, 45), Z, "VYK", "V-107", "A")

    # --- Hala Modřice: calculations, model, podklady ----------------------------
    def statika(ver, date, pages):
        lines = [f"Zakázka: {Z}   Akce: Hala Modřice   Verze výpočtu: {ver}   Datum: {date}",
                 f"Zpracoval: Ing. Jan Dvořák, {KANCELAR}", "",
                 "1. Popis konstrukce: ocelová hala 36 × 18 m, sloupy HEA 280, vazníky příhradové.",
                 "2. Zatížení: stálé, sníh, vítr (syntetický text pro workshop).",
                 "3. Posouzení sloupů – vyhovuje (ilustrační, neověřeno).",
                 "4. Posouzení vazníku – vyhovuje (ilustrační, neověřeno).",
                 "5. Kotvení – patky 500 × 500 mm, 4 × M24.",
                 "6. Závěr: text je smyšlený, nejde o skutečný statický výpočet."]
        return text_pdf(f"STATICKÝ VÝPOČET – {ver}", lines * 2, pages=pages)
    sv3 = statika("v3", "10.03.2026", 3)
    add(TOP + "Statika_Modrice_v3.pdf", sv3, T(2026, 3, 10, 12, 0), Z, "VYP", "SV", "v3")
    add(TOP + "stare/staticky vypocet hala v2.pdf", statika("v2", "25.02.2026", 3), T(2026, 2, 25, 16, 0), Z, "VYP", "SV", "v2",
        "stara_revize")
    add(TOP + "Stažené soubory/Statický výpočet - Modřice (1).pdf", sv3, T(2026, 3, 11, 9, 0), Z, "VYP", "SV", "v3",
        "duplicita", "stejný obsah jako Statika_Modrice_v3.pdf")
    add(TOP + "vypocet_patky_HM.xlsx", xlsx({"Patky": [
        ["Výpočet patek – Hala Modřice (Z2026-014)"], [],
        ["Patka", "Osa", "N_Ed [kN]", "Rozměr [mm]", "Kotvy"],
        *[[f"P{i + 1}", f"{'ABC'[i % 3]}/{i // 3 + 1}", 120 + 7 * i, "500×500", "4× M24"] for i in range(21)]]},
        "Výpočet patek"), T(2026, 3, 3, 10, 0), Z, "VYP", "patky")
    add(TOP + "Hala_Modrice_2026-03-10.ifc", ifc("Hala_Modrice.ifc", Z, "2026-03-10T14:22:00"), T(2026, 3, 10, 14, 22),
        Z, "MOD", "2026-03-10")
    add(TOP + "stare/Hala_Modrice_2026-02-20.ifc", ifc("Hala_Modrice.ifc", Z, "2026-02-20T10:05:00", 9), T(2026, 2, 20, 10, 5),
        Z, "MOD", "2026-02-20", stav="stara_revize")
    add(TOP + "Modrice_situace_architekt.dwg", dwg("Projekt: Hala Modrice - situace (architekt)"), T(2026, 2, 5, 9, 0),
        Z, "POD", "situace-architekt")
    add(TOP + "od Pavla/podklad_architekt_rez_A-A.dwg", dwg("Projekt: Hala Modrice - rez A-A (architekt)"), T(2026, 2, 5, 9, 2),
        Z, "POD", "rez-A-A-architekt", pozn="název neobsahuje zakázku; poznáte z obsahu/kontextu")

    # --- Hala Modřice: correspondence -------------------------------------------
    malá = "Ing. Petra Malá <petra.mala@stavby-novak.example>"
    my = "Ing. Karel Svoboda <svoboda@statika-detail.example>"
    add(TOP + "RE_ Hala Modřice - zvýšení haly.eml", eml(malá, my, "RE: Hala Modřice - zvýšení haly",
        datetime(2026, 3, 13, 8, 41, tzinfo=CET), (
        "Dobrý den, pane Svobodo,\n\npo kontrolním dni potvrzujeme požadavek na zvýšení světlé výšky haly "
        "o 300 mm (na 7,8 m).\nProsím o informaci, které výkresy se mění a kdy je dostaneme.\n\n"
        "S pozdravem\nPetra Malá\nStavby Novák s.r.o. (fiktivní)\n")), T(2026, 3, 13, 8, 41), Z, "KOR", "2026-03-13")
    add(TOP + "Fwd_ nabídka ocel HM.eml", eml(malá, my, "Fwd: nabídka ocel HM",
        datetime(2026, 3, 4, 15, 12, tzinfo=CET), (
        "Přeposílám nabídku dodavatele oceli pro halu Modřice, soubor nabidka_ocel_OcelServis_2026-03.xlsx.\n"
        "Prosím o kontrolu, zda sedí množství s vaším výkazem.\n\nPM\n")), T(2026, 3, 4, 15, 12), Z, "KOR", "2026-03-04")
    add(TOP + "pozvanka KD 12.3.eml", eml("Tomáš Král <kral@ocelmont.example>", my, "Pozvánka: kontrolní den Modřice 12. 3.",
        datetime(2026, 3, 9, 7, 30, tzinfo=CET), (
        "Zvu na kontrolní den č. 3 na stavbě haly v Modřicích, 12. 3. 2026 v 9:00.\n"
        "Témata: výška haly, kotvení, schodiště.\nT. Král, OcelMont (fiktivní)\n")), T(2026, 3, 9, 7, 30), Z, "KOR", "2026-03-09")
    add(TOP + "od Pavla/email od Pavla.txt", (
        "Od: Pavel Novotný (geodet)\nDatum: 2. 3. 2026\nVěc: Modřice – kotvy\n\n"
        "Ahoj, posílám opravený kotevní plán (rev B) podle zaměření. Starou revizi A prosím už nepoužívejte.\n"
        "Pavel\n").encode("utf-8"), T(2026, 3, 2, 17, 10), Z, "KOR", "2026-03-02")
    zapis = """ZÁPIS Z KONTROLNÍHO DNE č. 3
Akce: Hala Modřice (zak. Z2026-014)
Datum: 12. 3. 2026, místo: staveniště Modřice
Přítomni: Ing. Petra Malá (investor, Stavby Novák s.r.o. (fiktivní)), Ing. Jan Dvořák (statika),
Tomáš Král (montáž, OcelMont), Ing. Karel Svoboda (detailing)

1) Výška haly – investor požaduje zvýšení světlé výšky o 300 mm (na 7,8 m).
   -> ÚKOL: přepočítat sloupy a upravit V-102 a V-103 | zodp.: Dvořák, Svoboda | termín: 26. 3. 2026
2) Kotvení – geodet zaměřil patky, odchylky do 8 mm, bez připomínek.
3) Schodiště – změna šířky ramene na 1100 mm.
   -> ÚKOL: upravit výkres schodiště V-107 (rev B) | zodp.: Svoboda | termín: 20. 3. 2026
4) Fakturace – investor žádá rozpis fakturace po etapách.
   -> ÚKOL: poslat rozpis fakturace | zodp.: Svoboda | termín: 26. 3. 2026
5) Opláštění – dodavatel paždíků zatím není vybrán.
   -> ÚKOL: potvrdit dodavatele paždíků | zodp.: Malá | termín: 19. 3. 2026
6) Montáž – OcelMont potřebuje montážní schéma do konce měsíce.
   -> ÚKOL: připravit montážní schéma | zodp.: Svoboda, Král | termín: 31. 3. 2026

Příští KD: 26. 3. 2026 v 9:00
Zapsal: K. Svoboda
"""
    add(TOP + "zapis_KD_2026-03-12.txt", zapis.encode("utf-8"), T(2026, 3, 12, 14, 0), Z, "ZAP", "2026-03-12")
    poznamky = """14.3. volala Malá (investor Modřice)
- chce vědět, kdy budou upravené výkresy po zvýšení haly
- ptá se na cenu za změnu: detailing cca 16 h + statika cca 6 h
- sazba dle SoD 950 Kč/h bez DPH
- chce to mailem do pátku 20.3., kopie Dvořák
- pozor: V-103 už odešla v rev C 10.3. -> nová bude rev D, V-102 bude rev D
- termín úprav: 26.3. (jak je v zápisu z KD)
"""
    add(TOP + "poznamky_telefon.txt", poznamky.encode("utf-8"), T(2026, 3, 14, 11, 20), Z, "POZ", "2026-03-14")

    # --- Hala Modřice: commercial ----------------------------------------------
    steel = [["Nabídka – dodávka hutního materiálu"], ["Dodavatel:", "OcelServis s.r.o. (fiktivní)"],
             ["Akce:", "Hala Modřice"], ["Platnost:", "do 31. 3. 2026"], [],
             ["Profil", "Jakost", "Množství [t]", "Cena [Kč/t]", "Celkem [Kč]"]]
    for prof, q, p in [("HEA 280", 14.2, 31500), ("IPE 300", 9.8, 30900), ("L 70×7", 2.1, 33800),
                       ("UPE 160", 3.4, 32100), ("Plech P15", 1.6, 29800), ("Trubka 89×4", 0.9, 41200)]:
        steel.append([prof, "S355" if prof[0] in "HI" else "S235", q, p, round(q * p)])
    steel_b = xlsx({"Nabídka": steel}, "Nabídka oceli")
    add(TOP + "nabidka_ocel_OcelServis_2026-03.xlsx", steel_b, T(2026, 3, 4, 15, 20), Z, "NAB", "2026-03-04_OcelServis")
    add(TOP + "~$nabidka_ocel_OcelServis_2026-03.xlsx", b"\x0bKarel Svoboda" + bytes(149), T(2026, 3, 18, 10, 2),
        stav="smeti", pozn="zámek Excelu (otevřený soubor), nepatří k dokumentům")

    def det_offer(rows, ver):
        return xlsx({"Nabídka": [[f"Nabídka detailingu {ver} – Hala Modřice (Z2026-014)"],
                                 ["Klient:", "Stavby Novák s.r.o. (fiktivní)"], [],
                                 ["Položka", "Počet", "Hodin/ks", "Sazba [Kč/h]", "Celkem [Kč]"],
                                 *[[n, k, h, 950, k * h * 950] for n, k, h in rows]]}, f"Nabídka {ver}")
    base = [("3D model Tekla", 1, 80), ("Montážní výkresy", 7, 6), ("Výrobní výkresy dílců", 60, 1.2), ("Výkaz materiálu", 1, 6)]
    add(TOP + "nabidka_detailing_Modrice.xlsx", det_offer(base, "v1"), T(2026, 1, 28, 10, 0), Z, "NAB", "detailing", "v1",
        "stara_revize")
    add(TOP + "nabidka_detailing_Modrice_v2.xlsx", det_offer(base[:2] + [("Výrobní výkresy dílců", 72, 1.2)] + base[3:], "v2"),
        T(2026, 2, 3, 9, 0), Z, "NAB", "detailing", "v2")
    add(TOP + "FA_2026_0042.pdf", text_pdf("FAKTURA – DAŇOVÝ DOKLAD č. FA-2026-0042", [
        f"Dodavatel: {KANCELAR}", "Odběratel: Stavby Novák s.r.o. (fiktivní)", f"Zakázka: {Z} – Hala Modřice",
        "Datum vystavení: 06.03.2026   Splatnost: 20.03.2026", "",
        "Dílenská dokumentace – 1. etapa (model, kotevní plán)      62 000 Kč",
        "DPH 21 %                                                    13 020 Kč",
        "Celkem k úhradě                                             75 020 Kč"]), T(2026, 3, 6, 13, 0), Z, "FAK", "FA-2026-0042")
    add(TOP + "faktura 51 Modrice.pdf", text_pdf("FAKTURA č. 51/2026", [
        "Dodavatel: GeoMěření Novotný (fiktivní)", f"Odběratel: {KANCELAR}", "Akce: Hala Modřice – zaměření patek",
        "Datum vystavení: 04.03.2026", "", "Zaměření kotev a patek                          8 500 Kč",
        "Neplátce DPH", "Celkem k úhradě                                 8 500 Kč"],
        author="GeoMěření Novotný"), T(2026, 3, 4, 18, 0), Z, "FAK", "51-2026")
    add(TOP + "Smlouva_o_dilo_Modrice_navrh.docx", docx_contract(), T(2026, 2, 10, 9, 0), Z, "SML", "smlouva-o-dilo-navrh")

    # --- Hala Modřice: photos ---------------------------------------------------
    shots = [("101512", "patky osa A"), ("101547", "kotvy osa 4"), ("102003", "sloupy osa A"), ("102210", "staveniště")]
    for i, (hhmmss, what) in enumerate(shots):
        dt = datetime(2026, 3, 5, int(hhmmss[:2]), int(hhmmss[2:4]), int(hhmmss[4:]))
        b = photo(f"Modřice KD 5.3.\n{what}\nZ2026-014", dt, f"Hala Modřice Z2026-014 – {what}", 100 + i)
        add(TOP + f"fotky KD 5.3/IMG_20260305_{hhmmss}.jpg", b, T(2026, 3, 5, int(hhmmss[:2]), int(hhmmss[2:4])),
            Z, "FOT", f"2026-03-05_{hhmmss}")
        if i == 0:
            first = b
    add(TOP + "Stažené soubory/IMG_20260305_101512 (1).jpg", first, T(2026, 3, 6, 8, 0), Z, "FOT", "2026-03-05_101512",
        stav="duplicita", pozn="stejná fotka jako fotky KD 5.3/IMG_20260305_101512.jpg")

    # --- junk -------------------------------------------------------------------
    add(TOP + "Thumbs.db", bytes(rng.getrandbits(8) for _ in range(1024)), T(2026, 3, 12, 10, 4), stav="smeti",
        pozn="systémový soubor Windows (náhledy)")
    add(TOP + "fotky KD 5.3/Thumbs.db", bytes(rng.getrandbits(8) for _ in range(1024)), T(2026, 3, 5, 18, 0),
        stav="smeti", pozn="systémový soubor Windows (náhledy)")
    add(TOP + "Nový textový dokument.txt", b"", T(2026, 3, 2, 8, 0), stav="smeti", pozn="prázdný soubor")

    # --- not project related ----------------------------------------------------
    add(TOP + "Scan_20260318_0003.pdf", text_pdf("POTVRZENÍ O ŠKOLENÍ BOZP", [
        "Jméno: Ing. Karel Svoboda", "Školení: BOZP pro práci na staveništi, 18.03.2026",
        "Platnost: 2 roky", "Školitel: BezpečněBrno (fiktivní)"]), T(2026, 3, 18, 15, 0), "", "", "", "",
        "nezarazeno", "interní doklad kanceláře, k žádné zakázce nepatří")
    add(TOP + "cenik_2026.xlsx", xlsx({"Ceník": [["Ceník prací 2026 – interní"], [],
        ["Činnost", "Sazba [Kč/h]"], ["Detailing Tekla", 950], ["Statika", 1150], ["Kreslení úprav", 750]]}, "Ceník"),
        T(2026, 1, 5, 9, 0), "", "", "", "", "nezarazeno", "interní ceník kanceláře, k žádné zakázce nepatří")

    # --- Lávka Z2026-009 ----------------------------------------------------------
    L = "Z2026-009"
    lh = [("A", "15.01.2026", "První vydání"), ("B", "12.02.2026", "Zesílení horního pásu")]
    lav_a = drawing_pdf(L, "L-01", "Hlavní nosník – pohled", "A", lh[:1], "lavka", "A1", "1:50", "15.01.2026")
    lav_b = drawing_pdf(L, "L-01", "Hlavní nosník – pohled", "B", lh, "lavka", "A1", "1:50", "12.02.2026")
    add(TOP + "Lavka_L-01_revA.pdf", lav_a, T(2026, 1, 15, 10, 0), L, "VYK", "L-01", "A", "stara_revize")
    add(TOP + "Lavka_L-01_revB.pdf", lav_b, T(2026, 2, 12, 16, 0), L, "VYK", "L-01", "B")
    add(TOP + "Nová složka/L-02 zábradlí.pdf", drawing_pdf(L, "L-02", "Zábradlí", "A", [("A", "20.01.2026", "První vydání")],
        "oplasteni", "A3", "1:20", "20.01.2026"), T(2026, 1, 20, 11, 0), L, "VYK", "L-02", "A")
    add(TOP + "Stažené soubory/statika_lavka_v1.pdf", text_pdf("STATICKÝ VÝPOČET – v1", [
        f"Zakázka: {L}   Akce: Lávka přes Svitavu   Verze výpočtu: v1",
        "Ilustrační syntetický text, nejde o skutečný výpočet."]), T(2026, 1, 12, 9, 0), L, "VYP", "SV", "v1")
    add(TOP + "FA_2026_0038.pdf", text_pdf("FAKTURA – DAŇOVÝ DOKLAD č. FA-2026-0038", [
        f"Dodavatel: {KANCELAR}", "Odběratel: Město Svitávka (fiktivní)", f"Zakázka: {L} – Lávka přes Svitavu",
        "Datum vystavení: 27.02.2026", "", "Dokumentace lávky – 1. etapa       48 000 Kč",
        "DPH 21 %                            10 080 Kč", "Celkem k úhradě                     58 080 Kč"]),
        T(2026, 2, 27, 10, 0), L, "FAK", "FA-2026-0038")
    add(TOP + "dotaz_lavka.eml", eml("Bc. Jana Horáková <horakova@svitavka.example>", my, "Dotaz – lávka, nátěr zábradlí",
        datetime(2026, 3, 16, 13, 5, tzinfo=CET),
        "Dobrý den, jaký odstín nátěru je uvažován pro zábradlí lávky? Děkuji, J. Horáková\n"),
        T(2026, 3, 16, 13, 5), L, "KOR", "2026-03-16")
    add(TOP + "IMG_20260402_083015.jpg", photo("Lávka Svitávka\nopěra 1\nZ2026-009", datetime(2026, 4, 2, 8, 30, 15),
        "Lávka přes Svitavu Z2026-009 – opěra 1", 300), T(2026, 4, 2, 8, 30), L, "FOT", "2026-04-02_083015")

    # --- Parkovací dům Bystrc Z2025-031 ------------------------------------------
    P = "Z2025-031"
    ph = [("A", "10.09.2025", "První vydání"), ("B", "02.10.2025", "Změna výztuže"), ("C", "20.10.2025", "Úprava ozubů"),
          ("D", "14.11.2025", "Skutečné provedení")]
    add(TOP + "Nová složka/PD_Bystrc_P-21_D.pdf", drawing_pdf(P, "P-21", "Skladba prefabrikátů – 2. NP", "D", ph, "prefa", "A2",
        "1:100", "14.11.2025"), T(2025, 11, 14, 15, 0), P, "VYK", "P-21", "D")
    add(TOP + "nabidka_prefa_PD_Bystrc.xlsx", xlsx({"Nabídka": [["Nabídka – prefabrikáty, Parkovací dům Bystrc"], [],
        ["Dílec", "Počet", "Cena/ks [Kč]"], ["Sloup S1", 24, 38500], ["Průvlak P1", 18, 52000], ["Panel TT", 64, 41000]]},
        "Nabídka prefa"), T(2025, 8, 20, 9, 0), P, "NAB", "prefa")
    add(TOP + "FA_2025_0117.pdf", text_pdf("FAKTURA – DAŇOVÝ DOKLAD č. FA-2025-0117", [
        f"Dodavatel: {KANCELAR}", "Odběratel: Parking Bystrc a.s. (fiktivní)", f"Zakázka: {P} – Parkovací dům Bystrc",
        "Datum vystavení: 15.12.2025", "", "Výrobní dokumentace prefabrikátů – závěrečná faktura    120 000 Kč",
        "DPH 21 %                                                   25 200 Kč",
        "Celkem k úhradě                                           145 200 Kč"]), T(2025, 12, 15, 9, 0), P, "FAK", "FA-2025-0117")
    add(TOP + "Nová složka/foto_PD_Bystrc_montaz.jpg", photo("PD Bystrc\nmontáž TT panelů\nZ2025-031",
        datetime(2025, 10, 30, 11, 0), "Parkovací dům Bystrc Z2025-031 – montáž", 400), T(2025, 10, 30, 11, 0),
        P, "FOT", "2025-10-30_montaz")
    add(TOP + "PD_Bystrc_model.ifc", ifc("PD_Bystrc.ifc", P, "2025-11-14T16:00:00", 15), T(2025, 11, 14, 16, 0),
        P, "MOD", "2025-11-14")


def main():
    DATA.mkdir(parents=True, exist_ok=True)
    DEMO.mkdir(parents=True, exist_ok=True)
    build()
    order = sorted(range(len(FILES)), key=lambda i: FILES[i][0].lower())
    out = DATA / "zakazka_neporadek.zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in order:
            path, data, when = FILES[i]
            zi = zipfile.ZipInfo(path, when)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zf.writestr(zi, data)
    with open(DATA / "seznam_zakazek.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["cislo_zakazky", "nazev", "slozka", "klient", "kontakt", "misto", "typ_konstrukce", "stav",
                    "zodpovedny", "prezdivky_v_nazvech"])
        for k, z in ZAKAZKY.items():
            w.writerow([k, z["nazev"], z["slozka"], z["klient"], z["kontakt"], z["misto"], z["typ"], z["stav"],
                        z["zodpovedny"], z["prezdivky"]])
    with open(DEMO / "klic_zarazeni.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(KEY[0]), delimiter=";")
        w.writeheader()
        for i in order:
            w.writerow(KEY[i])
    size = out.stat().st_size
    print(f"{out.relative_to(ROOT)}: {len(FILES)} souborů, {size / 1024:.0f} kB")
    assert size < 2_000_000


if __name__ == "__main__":
    main()
