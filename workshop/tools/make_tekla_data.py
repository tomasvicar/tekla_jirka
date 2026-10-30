"""Syntetické exporty z Tekla Structures pro cvičení „Data z Tekly“.

Ocelová hala 24 × 48 m (9 rámů po 6 m, okap 6 m), dvě revize modelu (A, B).
Revize B obsahuje skutečné změny i úmyslné chyby k nalezení.
Vše je vymyšlené, plně deterministické (bez náhody, každé spuštění dá stejné soubory), žádná data klientů.

Spuštění (z kořene workshopu):
    uv run --with openpyxl python tools/make_tekla_data.py

Výstup: data/tekla/  (+ data/tekla_data.zip pro jedno stažení)
"""
from __future__ import annotations

import copy
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "tekla"
ZIP = ROOT / "data" / "tekla_data.zip"

PROJECT = "HALA SKLAD 24x48"
ORDER = "Z-2026-014"

# Profil: (kg/m, m2/m povrchu, kód DSTV, výška, šířka pásnice, tl. pásnice, tl. stojiny)
# Hodnoty odpovídají běžným tabulkám válcovaných profilů (zaokrouhleno).
PROFILES = {
    "HEA200": (42.3, 1.136, "I", 190, 200, 10.0, 6.5),
    "HEA300": (88.3, 1.717, "I", 290, 300, 14.0, 8.5),
    "IPE160": (15.8, 0.623, "I", 160, 82, 7.4, 5.0),
    "IPE400": (66.3, 1.467, "I", 400, 180, 13.5, 8.6),
    "UPE140": (14.5, 0.556, "U", 140, 65, 9.0, 5.0),
    "UPE160": (17.0, 0.616, "U", 160, 70, 9.5, 5.5),
    "UPE180": (19.7, 0.676, "U", 180, 75, 10.5, 5.5),
    "RHS100*100*5": (14.4, 0.383, "M", 100, 100, 5.0, 5.0),
    "RHS120*120*6": (20.7, 0.459, "M", 120, 120, 6.0, 6.0),
    "L70*7": (7.38, 0.273, "L", 70, 70, 7.0, 7.0),
    "L60*6": (5.42, 0.233, "L", 60, 60, 6.0, 6.0),
    "D20": (2.47, 0.063, "RO", 20, 20, 0, 0),
}
STEEL_DENSITY = 7850  # kg/m3


def profile_props(profile: str, length: float):
    """Vrátí (hmotnost kg/ks, povrch m2/ks)."""
    if profile.startswith("PL"):
        t, w = (float(x) for x in profile[2:].split("*"))
        weight = t * w * length * 1e-9 * STEEL_DENSITY
        area = 2 * w * length * 1e-6 + 2 * t * (w + length) * 1e-6
        return round(weight, 1), round(area, 3)
    kg_m, m2_m, *_ = PROFILES[profile]
    return round(kg_m * length / 1000, 1), round(m2_m * length / 1000, 3)


def P(pos, name, profile, length, material, cls, n=1, main=False):
    return dict(pos=pos, name=name, profile=profile, length=length, material=material,
                cls=cls, n=n, main=main)


def A(phase, pos, name, qty, finish, vykres, parts, dodavka="OCEL"):
    return dict(phase=phase, pos=pos, name=name, qty=qty, finish=finish,
                vykres=vykres, dodavka=dodavka, parts=parts)


# ---------------------------------------------------------------- revize A
REV_A = [
    # Fáze 1 – hlavní rámy
    A(1, "C1", "SLOUP RAMU", 18, "NATER_C3", "S-101", [
        P("1001", "SLOUP", "HEA300", 6150, "S355J2", 2, main=True),
        P("1002", "PATNI PLECH", "PL20*400", 450, "S355J2", 9),
        P("1003", "CELNI PLECH", "PL20*300", 700, "S355J2", 9),
        P("1004", "VYZTUHA", "PL10*140", 262, "S235JR", 9, n=4),
    ]),
    A(1, "B1", "PRICLE RAMU", 18, "NATER_C3", "S-102", [
        P("1101", "PRICLE", "IPE400", 11905, "S355J2", 3, main=True),
        P("1102", "NABEH", "IPE400", 1500, "S355J2", 3),
        P("1103", "CELNI PLECH", "PL20*300", 700, "S355J2", 9),
        P("1104", "VYZTUHA", "PL10*85", 373, "S235JR", 9, n=4),
        P("1105", "VRCHOLOVY PLECH", "PL20*220", 600, "S355J2", 9),
    ]),
    A(1, "C2", "STITOVY SLOUP", 4, "NATER_C3", "S-103", [
        P("1201", "SLOUP", "HEA200", 6480, "S355J2", 2, main=True),
        P("1202", "PATNI PLECH", "PL15*300", 300, "S235JR", 9),
        P("1203", "HLAVOVY PLECH", "PL10*200", 250, "S235JR", 9),
    ]),
    A(1, "C3", "STITOVY SLOUP", 2, "NATER_C3", "S-103", [
        P("1301", "SLOUP", "HEA200", 7000, "S355J2", 2, main=True),
        P("1202", "PATNI PLECH", "PL15*300", 300, "S235JR", 9),
        P("1203", "HLAVOVY PLECH", "PL10*200", 250, "S235JR", 9),
    ]),
    # Fáze 2 – střecha
    A(2, "B2", "VAZNICE", 112, "ZINEK", "S-201", [
        P("2001", "VAZNICE", "IPE160", 5990, "S235JR", 5, main=True),
    ]),
    A(2, "B3", "OKAPOVA VAZNICE", 16, "ZINEK", "S-201", [
        P("2101", "VAZNICE", "UPE160", 5990, "S235JR", 5, main=True),
    ]),
    A(2, "B4", "VZPERA", 16, "NATER_C3", "S-202", [
        P("2201", "VZPERA", "RHS100*100*5", 5700, "S355J2H", 4, main=True),
        P("2202", "KONCOVY PLECH", "PL10*150", 150, "S235JR", 9, n=2),
    ]),
    A(2, "V1", "STRESNI ZTUZIDLO", 16, "ZINEK", "S-203", [
        P("2301", "TAHLO", "D20", 8460, "S235JR", 4, main=True),
    ]),
    # Fáze 3 – stěny
    A(3, "B5", "PAZDIK", 48, "ZINEK", "S-301", [
        P("3001", "PAZDIK", "UPE140", 5990, "S235JR", 6, main=True),
    ]),
    A(3, "B6", "PAZDIK STITU", 24, "ZINEK", "S-301", [
        P("3101", "PAZDIK", "UPE140", 5790, "S235JR", 6, main=True),
    ]),
    A(3, "V2", "STENOVE ZTUZIDLO", 8, "ZINEK", "S-302", [
        P("3201", "DIAGONALA", "L70*7", 8400, "S235JR", 4, main=True),
        P("3202", "STYCNIK", "PL10*150", 200, "S235JR", 9, n=2),
    ]),
    A(3, "V3", "ZTUZIDLO STITU", 4, "ZINEK", "S-303", [
        P("3301", "DIAGONALA", "L60*6", 4800, "S235JR", 4, main=True),
        P("3302", "STYCNIK", "PL10*150", 200, "S235JR", 9, n=2),
    ]),
]


def find(rev, pos):
    return next(a for a in rev if a["pos"] == pos)


# ---------------------------------------------------------------- revize B
REV_B = copy.deepcopy(REV_A)
# Skutečné změny projektu
find(REV_B, "C1")["parts"][1]["profile"] = "PL25*400"     # silnější patní plech (změna hmotnosti)
find(REV_B, "C3")["parts"][0]["length"] = 7120            # delší prostřední štítový sloup
find(REV_B, "B2")["qty"] = 128                            # o řadu vaznic víc na každé straně
find(REV_B, "B3")["parts"][0]["profile"] = "UPE180"       # okapová vaznice větší profil
find(REV_B, "V2")["parts"][0]["material"] = "S355J2"      # stěnové ztužidlo jiná ocel
REV_B.remove(find(REV_B, "V3"))                           # ztužidlo štítu nahrazeno rámem vrat
b6 = find(REV_B, "B6")
b6["qty"] = 20                                            # paždíky štítu: 4 se zkrátí kvůli vratům
REV_B.append(A(3, "C4", "SLOUPEK VRAT", 2, "NATER_C3", "", [   # CHYBA: prázdné UDA_VYKRES
    P("3401", "SLOUPEK", "RHS120*120*6", 4200, "S355J2H", 2, main=True),
    P("3402", "PATNI PLECH", "PL15*250", 250, "S235JR", 9),
]))
# CHYBA: překlad vrat přidán po posledním číslování → pozice "?"
# CHYBA: chybí materiál u nového dílu překladu
REV_B.append(A(3, "?", "PREKLAD VRAT", 1, "NATER_C3", "S-304", [
    P("?", "PREKLAD", "RHS120*120*6", 4000, "S355J2H", 3, main=True),
    P("?", "KONCOVY PLECH", "PL10*150", 150, "", 9, n=2),
]))
# CHYBA: duplicitní značka – zkrácené paždíky štítu dostaly stejnou značku B6
REV_B.append(A(3, "B6", "PAZDIK STITU", 4, "ZINEK", "S-301", [
    P("3101", "PAZDIK", "UPE140", 2690, "S235JR", 6, main=True),
]))
# CHYBA: překlep v profilu štítového sloupu
find(REV_B, "C2")["parts"][0]["profile_text"] = "HAE200"
# CHYBA: prázdné UDA_DODAVKA u stěnového ztužidla
find(REV_B, "V2")["dodavka"] = ""
# CHYBA: chybí materiál u výztuhy příčle
find(REV_B, "B1")["parts"][3]["material"] = ""

PHASE_ORDER = {1: 0, 2: 1, 3: 2}


def fmt(x, nd=1):
    s = f"{x:.{nd}f}"
    return s.replace(".", ",")


def rows_for(rev):
    parts, asms = [], []
    for a in sorted(rev, key=lambda a: (PHASE_ORDER[a["phase"]], a["pos"] == "?", a["pos"])):
        single_w = single_area = 0.0
        main = next(p for p in a["parts"] if p["main"])
        for p in a["parts"]:
            w, ar = profile_props(p["profile"], p["length"])
            single_w += w * p["n"]
            single_area += ar * p["n"]
            parts.append(dict(PHASE=a["phase"], ASSEMBLY_POS=a["pos"], PART_POS=p["pos"],
                              QTY=a["qty"] * p["n"], NAME=p["name"],
                              PROFILE=p.get("profile_text", p["profile"]), MATERIAL=p["material"],
                              LENGTH=p["length"], WEIGHT=w, AREA=ar, CLASS=p["cls"],
                              FINISH=a["finish"], UDA_VYKRES=a["vykres"], UDA_DODAVKA=a["dodavka"]))
        asms.append(dict(PHASE=a["phase"], ASSEMBLY_POS=a["pos"], QTY=a["qty"], NAME=a["name"],
                         PROFILE=main.get("profile_text", main["profile"]), MATERIAL=main["material"],
                         LENGTH=main["length"], WEIGHT=round(single_w, 1),
                         WEIGHT_TOTAL=round(single_w, 1) * a["qty"], AREA=round(single_area, 2),
                         FINISH=a["finish"], UDA_VYKRES=a["vykres"], UDA_DODAVKA=a["dodavka"]))
    return parts, asms


def write_report(path, title, rev, date, cols, rows, total_line):
    lines = [
        "TEKLA STRUCTURES REPORT",
        f"Report:;{title}",
        f"Projekt:;{PROJECT};Zakazka:;{ORDER}",
        f"Model revize:;{rev};Datum exportu:;{date}",
        "Jednotky:;LENGTH [mm];WEIGHT [kg/ks];AREA [m2/ks]",
        "",
        ";".join(cols),
    ]
    for r in rows:
        vals = []
        for c in cols:
            v = r[c]
            if c in ("WEIGHT", "WEIGHT_TOTAL"):
                v = fmt(v, 1)
            elif c == "AREA":
                v = fmt(v, 3 if "PART_POS" in cols else 2)
            vals.append(str(v))
        lines.append(";".join(vals))
    lines += ["", total_line]
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8-sig")


def write_nc(folder, parts_b):
    """DSTV NC1 soubory – úmyslně „starší“ export: některé díly neodpovídají revizi B."""
    folder.mkdir(parents=True, exist_ok=True)
    nc = {}
    for r in parts_b:
        if r["PROFILE"].startswith("PL") or r["PROFILE"] == "D20" or r["PART_POS"] == "?":
            continue
        key = r["PART_POS"]
        if key in nc:  # stejná pozice dílu ve více sestavách → sečíst kusy
            nc[key]["QTY"] += r["QTY"]
            continue
        nc[key] = dict(r)
    nc.pop("3401")                                   # chybí NC pro nový sloupek vrat
    nc["2101"]["PROFILE"] = "UPE160"                 # NC z revize A: starý profil
    nc["1301"]["LENGTH"] = 7000                      # NC z revize A: stará délka
    nc["1201"]["PROFILE"] = "HEA200"                 # NC má správný profil (překlep je jen v reportu)
    nc["3301"] = dict(PART_POS="3301", PROFILE="L60*6", MATERIAL="S235JR", LENGTH=4800,
                      QTY=4, PHASE=3)                # zbylý soubor po odstraněném ztužidle
    for key, r in sorted(nc.items()):
        kg_m, m2_m, code, h, b, tf, tw = PROFILES[r["PROFILE"]]
        holes = []
        if code == "I" and r["LENGTH"] > 3000:
            for x in (60, r["LENGTH"] - 60):
                for y in (h / 2 - 40, h / 2 + 40):
                    holes.append(f"  v  {x:8.2f}s {y:8.2f}  22.00")
        mat = r["MATERIAL"] or "S235JR"
        txt = [
            "ST",
            "** Tekla Structures - synteticky export pro workshop",
            f"  {ORDER}",
            f"  S-{key}",
            f"  {r['PHASE']}",
            f"  {key}",
            f"  {mat}",
            f"  {r['QTY']}",
            f"  {r['PROFILE']}",
            f"  {code}",
            f"  {r['LENGTH']:8.2f}",
            f"  {r['LENGTH']:8.2f}",
            f"  {h:8.2f}",
            f"  {b:8.2f}",
            f"  {tf:8.2f}",
            f"  {tw:8.2f}",
            "  0.00",
            f"  {kg_m:8.3f}",
            f"  {m2_m:8.3f}",
            "  0.000", "  0.000", "  0.000", "  0.000",
            "  -", "  -", "  -", "  -",
        ]
        if holes:
            txt += ["BO"] + holes
        txt.append("EN")
        (folder / f"{key}.nc1").write_text("\n".join(txt) + "\n", encoding="ascii")
    (folder / "dstv_nc.log").write_text(
        "Tekla Structures NC export\nDatum: 2026-10-20 14:02\nModel: HALA_24x48\n"
        f"Exportovano souboru: {len(nc)}\nPreskoceno (plechy a tahla nejsou v tomto exportu): ano\n",
        encoding="ascii")


def write_pricelist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Material"
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="DDEBF7")
    ws.append(["Cenik oceli (synteticke ceny, bez DPH, platnost 11/2026)"])
    ws.append([])
    ws.append(["Typ profilu", "Popis", "S235 [Kc/kg]", "S355 [Kc/kg]"])
    for row in [
        ("HEA", "valcovane H (HEA, HEB)", 32.5, 34.0),
        ("IPE", "valcovane I", 31.5, 33.0),
        ("UPE", "valcovane U", 33.0, 34.5),
        ("RHS", "uzavrene profily (RHS, SHS, CHS)", 41.0, 43.5),
        ("L", "uhelniky", 34.0, 35.5),
        ("D", "kruhove tyce (tahla)", 38.0, 40.0),
        ("PL", "plechy", 29.0, 30.5),
    ]:
        ws.append(list(row))
    ws.append([])
    ws.append(["Typ profilu = zacatek oznaceni profilu (HEA300 -> HEA, PL20*400 -> PL)."])
    ws.append(["Jakost: S235JR -> S235; S355J2, S355J2H -> S355."])
    for c in ws[3]:
        c.font, c.fill = bold, head
    ws["A1"].font = Font(bold=True, size=12)

    ws = wb.create_sheet("Povrch")
    ws.append(["Povrchova uprava (podle sloupce FINISH)"])
    ws.append([])
    ws.append(["FINISH", "Popis", "Cena", "Jednotka"])
    ws.append(["NATER_C3", "otryskani Sa 2,5 + naterovy system C3, 160 um", 420, "Kc/m2"])
    ws.append(["ZINEK", "zarove zinkovani ponorem", 19, "Kc/kg"])
    for c in ws[3]:
        c.font, c.fill = bold, head
    ws["A1"].font = Font(bold=True, size=12)

    ws = wb.create_sheet("Prace a ostatni")
    ws.append(["Polozka", "Hodnota", "Jednotka", "Poznamka"])
    for row in [
        ("Prorez materialu", 5, "%", "prirazka k cene materialu"),
        ("Dilenska vyroba", 26, "Kc/kg", "z hmotnosti konstrukce"),
        ("Montaz", 14, "Kc/kg", "z hmotnosti konstrukce"),
        ("Spojovaci material", 3, "%", "z ceny materialu vcetne prorezu"),
        ("Doprava", 9500, "Kc/jizda", "1 jizda = max. 20 t"),
        ("DPH", 21, "%", ""),
    ]:
        ws.append(list(row))
    for c in ws[1]:
        c.font, c.fill = bold, head
    for sh in wb.worksheets:
        for col, width in zip("ABCD", (22, 48, 14, 28)):
            sh.column_dimensions[col].width = width
    wb.save(path)


README = """DATA Z TEKLY - SYNTETICKA CVICNA DATA
=====================================
Vse je vymyslene (generator tools/make_tekla_data.py, deterministicky).
Ocelova hala 24 x 48 m, 9 ramu po 6 m, vyska okapu 6 m.

hala_revA_assembly_list.csv  seznam sestav, revize A (12. 10. 2026)
hala_revA_part_list.csv      seznam dilu,   revize A
hala_revB_assembly_list.csv  seznam sestav, revize B (27. 10. 2026)
hala_revB_part_list.csv      seznam dilu,   revize B
cenik.xlsx                   cenik oceli, povrchove upravy a prace
nc_dstv/                     DSTV (NC1) soubory pro dilenske stroje + dstv_nc.log

Format CSV: kodovani UTF-8, oddelovac strednik, desetinna carka.
Nad tabulkou je 6 radku hlavicky reportu, pod tabulkou prazdny radek a radek CELKEM.
WEIGHT a AREA jsou na 1 kus. QTY je pocet kusu v celem modelu.
Fáze (PHASE): 1 = hlavni ramy, 2 = strecha, 3 = steny.
CLASS: 2 sloupy, 3 nosniky, 4 ztuzidla, 5 vaznice, 6 pazdiky, 9 plechy.
UDA_VYKRES = cislo dilenskeho vykresu, UDA_DODAVKA = kdo dodava.
"""


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    part_cols = ["PHASE", "ASSEMBLY_POS", "PART_POS", "QTY", "NAME", "PROFILE", "MATERIAL",
                 "LENGTH", "WEIGHT", "AREA", "CLASS", "FINISH", "UDA_VYKRES", "UDA_DODAVKA"]
    asm_cols = ["PHASE", "ASSEMBLY_POS", "QTY", "NAME", "PROFILE", "MATERIAL", "LENGTH",
                "WEIGHT", "WEIGHT_TOTAL", "AREA", "FINISH", "UDA_VYKRES", "UDA_DODAVKA"]
    for tag, rev, date in [("A", REV_A, "2026-10-12"), ("B", REV_B, "2026-10-27")]:
        parts, asms = rows_for(rev)
        tot_p = sum(r["QTY"] * r["WEIGHT"] for r in parts)
        tot_a = sum(r["WEIGHT_TOTAL"] for r in asms)
        write_report(OUT / f"hala_rev{tag}_part_list.csv", "Part_list (seznam dilu)", tag, date,
                     part_cols, parts,
                     f"CELKEM;;;{sum(r['QTY'] for r in parts)};;;;;{fmt(tot_p)};;;;;")
        write_report(OUT / f"hala_rev{tag}_assembly_list.csv", "Assembly_list (seznam sestav)", tag,
                     date, asm_cols, asms,
                     f"CELKEM;;{sum(r['QTY'] for r in asms)};;;;;;{fmt(tot_a)};;;;")
        print(f"rev {tag}: {len(parts)} radku dilu, {len(asms)} radku sestav, "
              f"{tot_p:.1f} kg (dily) / {tot_a:.1f} kg (sestavy)")
        if tag == "B":
            write_nc(OUT / "nc_dstv", parts)
    write_pricelist(OUT / "cenik.xlsx")
    (OUT / "README.txt").write_text(README, encoding="utf-8")
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(OUT.rglob("*")):
            if f.is_file():
                info = zipfile.ZipInfo(str(f.relative_to(OUT.parent)), date_time=(2026, 10, 30, 12, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(info, f.read_bytes())
    print("hotovo:", OUT, ZIP)


if __name__ == "__main__":
    main()
