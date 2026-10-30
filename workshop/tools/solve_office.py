"""Referenční řešení cvičení Kancelář (kancelar.qmd).

Spuštění (z kořene workshop/, nejdřív tools/make_office_data.py):
    uv run --with pypdf --with openpyxl --with pillow python tools/solve_office.py

Pracuje jen na kopii v dočasné složce; data/kancelar/ nemění. Výsledky do demos/kancelar/:
    inventura.csv, inventura_souhrn.txt          krok 1 – co ve složce je
    plan_zarazeni.csv                            krok 2 – návrh (co kam), ke schválení
    prejmenovani_log.csv, struktura.txt          krok 3 – provedeno na kopii + záznam pro vrácení
    Z2026-014_predavaci_protokol.xlsx,           krok 4 – seznam výkresů a jedno PDF
    Z2026-014_vykresy_platne.pdf
    struktura_nova_zakazka.txt                   krok 5 – test skillu nova-zakazka
    ukoly_z_KD.xlsx, email_klientovi.txt         krok 6 – bonus
Na konci porovná zařazení s klíčem demos/kancelar/klic_zarazeni.csv.
"""
from __future__ import annotations

import csv
import os
import hashlib
import importlib.util
import io
import re
import shutil
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "kancelar"
OUT = ROOT / "demos" / "kancelar"
SKILL = OUT / "nova-zakazka" / "nova_zakazka.py"
DNES = date(2026, 3, 20)  # pevné datum → reprodukovatelné výstupy

spec = importlib.util.spec_from_file_location("nova_zakazka", SKILL)
nz = importlib.util.module_from_spec(spec)
spec.loader.exec_module(nz)

SLOZKA_TYPU = {"POD": "01_Podklady", "VYP": "02_Vypocty", "MOD": "03_Model", "VYK": "04_Vykresy",
               "KOR": "05_Korespondence", "ZAP": "05_Korespondence", "POZ": "05_Korespondence",
               "NAB": "06_Nabidky_a_fakturace", "SML": "06_Nabidky_a_fakturace", "FAK": "06_Nabidky_a_fakturace",
               "FOT": "07_Fotodokumentace"}
SMETI = {"thumbs.db", "desktop.ini", ".ds_store"}


def ascii_slug(s: str) -> str:
    return nz.bez_diakritiky(s).replace("_", "-").strip("-")


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ------------------------------------------------------------------ reading content
def read_text(p: Path) -> tuple[str, dict]:
    """Return (text, metadata) for any file type in the folder."""
    ext = p.suffix.lower()
    meta: dict = {}
    try:
        if ext == ".pdf":
            r = PdfReader(p)
            meta["stran"] = len(r.pages)
            return "\n".join(pg.extract_text() or "" for pg in r.pages), meta
        if ext == ".eml":
            m = BytesParser(policy=policy.default).parsebytes(p.read_bytes())
            meta["datum"] = parsedate_to_datetime(m["Date"]).date()
            return f"{m['Subject']}\n{m['From']}\n{m.get_content()}", meta
        if ext == ".txt":
            return p.read_text(encoding="utf-8", errors="replace"), meta
        if ext == ".xlsx":
            wb = load_workbook(p, read_only=True, data_only=True)
            cells = [str(v) for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v is not None]
            return "\n".join(cells), meta
        if ext == ".docx":
            xml = zipfile.ZipFile(p).read("word/document.xml").decode("utf-8")
            return re.sub(r"<[^>]+>", " ", xml), meta
        if ext == ".ifc":
            return p.read_text(encoding="ascii", errors="replace"), meta
        if ext == ".dwg":
            return p.read_bytes().decode("ascii", errors="ignore"), meta
        if ext in (".jpg", ".jpeg", ".png"):
            ex = Image.open(p).getexif()
            meta["datum_foto"] = ex.get(0x0132, "")
            return str(ex.get(0x010E, "")), meta
    except Exception as e:  # noqa: BLE001 – a damaged file is reported, not fatal
        meta["chyba"] = str(e)
    return "", meta


def load_zakazky():
    rows = list(csv.DictReader((DATA / "seznam_zakazek.csv").open(encoding="utf-8-sig"), delimiter=";"))
    return {r["cislo_zakazky"]: r for r in rows}


def find_zakazka(text: str, path: str, zak: dict) -> tuple[str, str]:
    nums = set(re.findall(r"Z\d{4}-\d{3}", text)) & set(zak)
    if len(nums) == 1:
        return nums.pop(), "číslo zakázky v obsahu"
    hits = set()
    for k, r in zak.items():
        for alias in [a.strip() for a in r["prezdivky_v_nazvech"].split(";")]:
            if re.search(rf"(?<![\w]){re.escape(alias)}(?![\w])", text + "\n" + path, re.I):
                hits.add(k)
    if len(hits) == 1:
        return hits.pop(), "název akce v obsahu/názvu"
    return "", "nejasné" if hits else "žádná stopa"


def classify(p: Path, rel: str, text: str, meta: dict) -> dict:
    ext = p.suffix.lower()
    d = dict(typ="", cislo="", revize="", nazev="", datum="", meritko="", format="")
    first = lambda pat, s=text, fl=re.M: (m.group(1).strip() if (m := re.search(pat, s, fl)) else "")  # noqa: E731
    if ext == ".pdf":
        if "ČÍSLO VÝKRESU" in text:
            d.update(typ="VYK", cislo=first(r"^ČÍSLO VÝKRESU\n([A-Z]-\d+)"), revize=first(r"^REVIZE\n([A-Z])$"),
                     nazev=first(r"^VÝKRES\n(.+)$"), datum=first(r"^DATUM\n(.+)$"))
            mf = first(r"^MĚŘÍTKO / FORMÁT\n(.+)$")
            d["meritko"], d["format"] = [s.strip() for s in mf.split("/")]
        elif "STATICKÝ VÝPOČET" in text:
            d.update(typ="VYP", cislo="SV", revize=first(r"Verze výpočtu: (v\d+)"))
        elif "FAKTURA" in text:
            d.update(typ="FAK", cislo=first(r"FAKTURA.*?č\. (\S+)").replace("/", "-"))
    elif ext == ".eml":
        d.update(typ="KOR", cislo=f"{meta['datum']:%Y-%m-%d}")
    elif ext == ".txt":
        if "ZÁPIS" in text:
            dd, mm, yy = re.search(r"Datum: (\d+)\. (\d+)\. (\d{4})", text).groups()
            d.update(typ="ZAP", cislo=f"{yy}-{int(mm):02d}-{int(dd):02d}_KD")
        elif re.match(r"Od: ", text):
            dd, mm, yy = re.search(r"Datum: (\d+)\. (\d+)\. (\d{4})", text).groups()
            d.update(typ="KOR", cislo=f"{yy}-{int(mm):02d}-{int(dd):02d}")
        else:
            d.update(typ="POZ", cislo=f"{datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d}")
    elif ext == ".xlsx":
        head = text.split("\n", 1)[0]
        if head.lower().startswith("výpočet"):
            d.update(typ="VYP", cislo=ascii_slug(head.split("–")[0].replace("Výpočet", "")).lower() or "tabulka")
        elif "nabídka" in head.lower():
            sup = first(r"Dodavatel:\n(\S+)")
            ver = first(r"Nabídka \w+ (v\d+)", head)
            kind = "detailing" if "detailing" in head else ("prefa" if "prefabrik" in head else "material")
            d.update(typ="NAB", cislo=kind + (f"_{ascii_slug(sup)}" if sup else ""), revize=ver)
    elif ext == ".docx" and "SMLOUVA" in text:
        d.update(typ="SML", cislo="smlouva-o-dilo" + ("_navrh" if "NÁVRH" in text else ""))
    elif ext == ".ifc":
        d.update(typ="MOD", cislo=first(r"FILE_NAME\('[^']*','(\d{4}-\d\d-\d\d)"))
    elif ext == ".dwg":
        stem = re.sub(r"(?i)modrice|hala|_architekt|podklad", "", p.stem)
        d.update(typ="POD", cislo=ascii_slug(stem).lower() + "_architekt")
    elif ext in (".jpg", ".jpeg"):
        dt = meta.get("datum_foto", "")
        d.update(typ="FOT", cislo=dt[:10].replace(":", "-") + "_" + dt[11:].replace(":", "") if dt else p.stem)
    return d


def revize_key(rev: str):
    return (int(rev[1:]) if rev[:1] == "v" and rev[1:].isdigit() else ord(rev[0])) if rev else 0


# ------------------------------------------------------------------ main
def main():
    OUT.mkdir(parents=True, exist_ok=True)
    zak = load_zakazky()
    tmp = Path(tempfile.mkdtemp(prefix="kancelar_"))
    try:
        src = tmp / "neporadek"
        with zipfile.ZipFile(DATA / "zakazka_neporadek.zip") as zf:
            zf.extractall(src)
            for info in zf.infolist():  # zachovat datum souborů jako Windows Průzkumník
                ts = datetime(*info.date_time).timestamp()
                (src / info.filename).touch()
                os.utime(src / info.filename, (ts, ts))
        base = src / "zakazka_neporadek"
        files = sorted([p for p in base.rglob("*") if p.is_file()], key=lambda p: str(p).lower())

        # ---------------- step 1: inventory (read only)
        inv = []
        for p in files:
            rel = p.relative_to(base).as_posix()
            text, meta = read_text(p)
            inv.append(dict(p=p, rel=rel, text=text, meta=meta, sha=sha(p), size=p.stat().st_size,
                            mtime=datetime.fromtimestamp(p.stat().st_mtime)))
        by_hash = defaultdict(list)
        for f in inv:
            if f["size"]:
                by_hash[f["sha"]].append(f)
        dup_groups = {h: g for h, g in by_hash.items() if len(g) > 1}
        gid = {h: i + 1 for i, h in enumerate(sorted(dup_groups, key=lambda h: dup_groups[h][0]["rel"].lower()))}
        with (OUT / "inventura.csv").open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["cesta", "pripona", "velikost_kB", "zmeneno", "sha256_zacatek", "skupina_duplicit"])
            for f in inv:
                w.writerow([f["rel"], f["p"].suffix.lower() or "(bez)", round(f["size"] / 1024, 1),
                            f"{f['mtime']:%Y-%m-%d %H:%M}", f["sha"][:12], gid.get(f["sha"], "")])

        # ---------------- step 2: classification and plan
        for f in inv:
            name = f["p"].name.lower()
            f.update(classify(f["p"], f["rel"], f["text"], f["meta"]))
            if name in SMETI or name.startswith("~$") or f["size"] == 0:
                f.update(zakazka="", stav="smeti", duvod="systémový/dočasný/prázdný soubor")
                continue
            f["zakazka"], f["duvod"] = find_zakazka(f["text"], f["rel"], zak)
            f["stav"] = "aktualni" if f["zakazka"] and f["typ"] else "nezarazeno"
        # duplicates: keep the oldest file of each group, others are copies
        for h, g in dup_groups.items():
            g = sorted(g, key=lambda f: f["mtime"])
            for f in g[1:]:
                if f["stav"] == "aktualni":
                    f["stav"], f["duvod"] = "duplicita", f"stejný obsah jako {g[0]['rel']}"
                    f["vzor"] = g[0]
        # superseded revisions / versions
        groups = defaultdict(list)
        for f in inv:
            if f["stav"] == "aktualni":
                key = (f["zakazka"], f["typ"]) if f["typ"] == "MOD" else (f["zakazka"], f["typ"], f["cislo"])
                if f["typ"] in ("VYK", "VYP", "NAB", "MOD"):
                    groups[key].append(f)
        for key, g in groups.items():
            if len(g) > 1:
                g.sort(key=lambda f: (revize_key(f["revize"]), f["cislo"]))
                for f in g[:-1]:
                    f["stav"], f["duvod"] = "stara_revize", f"nahrazeno: {g[-1]['rel']}"
        # filename vs title-block mismatch
        varovani = []
        for f in inv:
            if f["typ"] == "VYK":
                m = re.search(r"rev([A-Z])\b", f["p"].stem) or re.search(r"_([A-Z])$", f["p"].stem)
                if m and m.group(1) != f["revize"]:
                    varovani.append(f"{f['rel']}: v názvu rev {m.group(1)}, v razítku rev {f['revize']}")

        # target names
        used = set()
        for f in sorted(inv, key=lambda f: (f["stav"] != "aktualni", f["rel"].lower())):
            ext = f["p"].suffix.lower()
            if f["stav"] == "smeti":
                f["novy"], f["cil"], f["akce"] = "", "", "nekopírovat (smetí)"
                continue
            if f["stav"] == "nezarazeno":
                f["novy"] = f["p"].name
                f["cil"] = f"_K_ROZTRIDENI/{f['novy']}"
                f["akce"] = "kopírovat k ruční kontrole"
                continue
            novy = f"{f['zakazka']}_{f['typ']}_{f['cislo']}" + (f"_{f['revize']}" if f["revize"] else "")
            folder = {"aktualni": SLOZKA_TYPU[f["typ"]], "stara_revize": "99_Archiv/stare_revize",
                      "duplicita": "99_Archiv/duplicity"}[f["stav"]]
            cil = f"{zak[f['zakazka']]['slozka']}/{folder}/{novy}{ext}"
            n = 2
            while cil in used:
                cil = f"{zak[f['zakazka']]['slozka']}/{folder}/{novy}_kopie{n}{ext}"
                n += 1
            used.add(cil)
            f["novy"], f["cil"] = Path(cil).name, cil
            f["akce"] = {"aktualni": "kopírovat + přejmenovat", "stara_revize": "do archivu (stará revize)",
                         "duplicita": "do archivu (duplicita)"}[f["stav"]]

        cols = ["puvodni_cesta", "zakazka", "typ", "cislo", "revize", "stav", "akce", "nova_cesta", "duvod"]
        with (OUT / "plan_zarazeni.csv").open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(cols)
            for f in inv:
                w.writerow([f["rel"], f["zakazka"], f["typ"], f["cislo"], f["revize"], f["stav"], f["akce"],
                            f["cil"], f["duvod"]])

        # ---------------- step 3: execute on a copy (skeleton via the nova-zakazka script)
        dst = tmp / "Zakazky"
        dst.mkdir()
        for k, r in zak.items():
            created = nz.zaloz(k, r["nazev"], r["klient"], dst, dnes=DNES)
            assert created.name == r["slozka"], (created.name, r["slozka"])
        log = []
        for f in inv:
            if not f["cil"]:
                log.append([f["rel"], "", "vynechano", f["sha"][:12]])
                continue
            t = dst / f["cil"]
            t.parent.mkdir(parents=True, exist_ok=True)
            assert not t.exists()
            shutil.copy2(f["p"], t)
            assert sha(t) == f["sha"]
            log.append([f["rel"], f["cil"], "zkopirovano", f["sha"][:12]])
        with (OUT / "prejmenovani_log.csv").open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["puvodni_cesta", "nova_cesta", "akce", "sha256_zacatek"])
            w.writerows(log)

        def tree(root: Path) -> str:
            lines = [root.name + "/"]
            def walk(d: Path, pre: str):
                items = sorted(d.iterdir(), key=lambda x: (x.is_file(), x.name.lower()))
                for i, it in enumerate(items):
                    last = i == len(items) - 1
                    lines.append(pre + ("└── " if last else "├── ") + it.name + ("/" if it.is_dir() else ""))
                    if it.is_dir():
                        walk(it, pre + ("    " if last else "│   "))
            walk(root, "")
            return "\n".join(lines) + "\n"
        (OUT / "struktura.txt").write_text(tree(dst), encoding="utf-8")
        # summary of steps 1–3
        cnt_ext = Counter(f["p"].suffix.lower() or "(bez)" for f in inv)
        cnt_stav = Counter(f["stav"] for f in inv)
        cnt_zak = Counter(f["zakazka"] or "(nezařazeno/smetí)" for f in inv)
        souhrn = [f"Souborů celkem: {len(inv)} (v hlavní složce a {len({Path(f['rel']).parent for f in inv}) - 1} podsložkách)",
                  "Podle přípony: " + ", ".join(f"{k} {v}" for k, v in sorted(cnt_ext.items())),
                  f"Skupin stejných souborů (duplicit): {len(dup_groups)}, nadbytečných kopií: "
                  f"{sum(len(g) - 1 for g in dup_groups.values())}",
                  "Podle zakázky: " + ", ".join(f"{k} {v}" for k, v in sorted(cnt_zak.items())),
                  "Podle stavu: " + ", ".join(f"{k} {v}" for k, v in sorted(cnt_stav.items())),
                  "", "Rozpor název × razítko:", *[f"  - {v}" for v in varovani],
                  "", "Skupiny duplicit:"]
        for h in sorted(dup_groups, key=gid.get):
            souhrn.append(f"  {gid[h]}: " + " = ".join(f["rel"] for f in dup_groups[h]))
        (OUT / "inventura_souhrn.txt").write_text("\n".join(souhrn) + "\n", encoding="utf-8")

        # ---------------- step 4: transmittal + merged PDF for Z2026-014
        Z = "Z2026-014"
        vyk = sorted([f for f in inv if f["zakazka"] == Z and f["typ"] == "VYK" and f["stav"] == "aktualni"],
                     key=lambda f: f["cislo"])
        writer = PdfWriter()
        for f in vyk:
            start = len(writer.pages)
            writer.append(str(dst / f["cil"]))
            writer.add_outline_item(f"{f['cislo']} rev {f['revize']} – {f['nazev']}", start)
        writer.add_metadata({"/Title": f"{Z} – platné výkresy k {DNES:%d.%m.%Y}", "/Producer": "solve_office.py"})
        with (OUT / f"{Z}_vykresy_platne.pdf").open("wb") as fh:
            writer.write(fh)
        wb = Workbook()
        ws = wb.active
        ws.title = "Předávací protokol"
        r = zak[Z]
        bold, thin = Font(bold=True), Side(style="thin", color="888888")
        ws.append(["PŘEDÁVACÍ PROTOKOL – VÝKRESOVÁ DOKUMENTACE"])
        ws["A1"].font = Font(bold=True, size=14)
        for k, v in [("Zakázka", f"{Z} – {r['nazev']}"), ("Investor", r["klient"]),
                     ("Předává", "STATIKA & DETAIL s.r.o. (fiktivní)"), ("Datum předání", DNES.strftime("%d.%m.%Y"))]:
            ws.append([k, v])
            ws.cell(ws.max_row, 1).font = bold
        ws.append([])
        head = ["Poř.", "Číslo výkresu", "Název výkresu", "Revize", "Datum revize", "Měřítko", "Formát",
                "Počet listů", "Soubor"]
        ws.append(head)
        hr = ws.max_row
        for c in ws[hr]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="12355B")
        for i, f in enumerate(vyk, 1):
            ws.append([i, f["cislo"], f["nazev"], f["revize"], f["datum"], f["meritko"], f["format"],
                       f["meta"]["stran"], f["novy"]])
        for row in ws.iter_rows(min_row=hr, max_row=ws.max_row, max_col=len(head)):
            for c in row:
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                c.alignment = Alignment(vertical="center")
        ws.append([])
        ws.append(["Celkem výkresů", len(vyk), "", "Celkem listů", "", "", "", sum(f["meta"]["stran"] for f in vyk)])
        ws.cell(ws.max_row, 1).font = bold
        ws.append([])
        ws.append(["Převzal (jméno, podpis, datum):", "", "", "", "Předal:"])
        for col, wdt in zip("ABCDEFGHI", [8, 15, 32, 8, 13, 9, 8, 11, 30]):
            ws.column_dimensions[col].width = wdt
        ws.column_dimensions["A"].width = 16
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.properties.created = wb.properties.modified = datetime(2026, 3, 20, 8, 0)
        b = io.BytesIO()
        wb.save(b)
        (OUT / f"{Z}_predavaci_protokol.xlsx").write_bytes(_fix_zip(b.getvalue()))

        # ---------------- step 5: test the skill script on a new zakázka
        test = tmp / "test_skill"
        seznam = test / "seznam_zakazek.csv"
        test.mkdir()
        shutil.copy(DATA / "seznam_zakazek.csv", seznam)
        new = nz.zaloz("Z2026-015", "Sklad Kuřim", "Logistika Kuřim a.s. (fiktivní)", test, seznam, dnes=DNES)
        try:
            nz.zaloz("Z2026-015", "Sklad Kuřim", "", test)
            raise AssertionError("skript měl odmítnout existující složku")
        except FileExistsError:
            pass
        assert sum(1 for row in seznam.open(encoding="utf-8-sig") if row.startswith("Z2026-015;")) == 1
        (OUT / "struktura_nova_zakazka.txt").write_text(
            tree(new) + "\n--- README.txt ---\n" + (new / "README.txt").read_text(encoding="utf-8"), encoding="utf-8")

        # ---------------- step 6: tasks from the meeting minutes + e-mail draft
        zapis = next(f for f in inv if f["typ"] == "ZAP" and f["stav"] == "aktualni")["text"]
        ukoly = []
        bod = ""
        for line in zapis.splitlines():
            if m := re.match(r"(\d+)\) (.+?) – (.+)", line):
                bod = f"{m.group(1)}) {m.group(2)}"
            if m := re.search(r"ÚKOL: (.+?) \| zodp\.: (.+?) \| termín: (.+)$", line):
                ukoly.append([len(ukoly) + 1, bod, m.group(1), m.group(2), m.group(3), "otevřený"])
        wb = Workbook()
        ws = wb.active
        ws.title = "Úkoly z KD 12.3.2026"
        ws.append(["Č.", "Bod zápisu", "Úkol", "Zodpovídá", "Termín", "Stav"])
        for c in ws[1]:
            c.font = bold
        for u in ukoly:
            ws.append(u)
        for col, wdt in zip("ABCDEF", [5, 22, 48, 18, 14, 10]):
            ws.column_dimensions[col].width = wdt
        wb.properties.created = wb.properties.modified = datetime(2026, 3, 20, 8, 0)
        b = io.BytesIO()
        wb.save(b)
        (OUT / "ukoly_z_KD.xlsx").write_bytes(_fix_zip(b.getvalue()))

        pozn = next(f for f in inv if f["typ"] == "POZ" and f["stav"] == "aktualni")["text"]
        hod = [int(h) for h in re.findall(r"cca (\d+) h", pozn)]
        sazba = int(re.search(r"(\d+) Kč/h", pozn).group(1))
        cena = sum(hod) * sazba
        cena_txt = f"{cena:,}".replace(",", " ")
        (OUT / "email_klientovi.txt").write_text(f"""Komu: Ing. Petra Malá <petra.mala@stavby-novak.example>
Kopie: Ing. Jan Dvořák
Předmět: Hala Modřice (Z2026-014) – zvýšení haly: termín a cena úprav

Dobrý den, paní inženýrko,

děkuji za telefonát. K požadavku na zvýšení světlé výšky haly o 300 mm (na 7,8 m) posílám přehled:

Které výkresy se mění
- V-102 Sloupy a ztužidla – osa A: nově revize D (dosud platná C)
- V-103 Střešní vazník: nově revize D (dosud platná C, odeslaná 10. 3. 2026)
- před úpravou výkresů proběhne přepočet sloupů (Ing. Dvořák)

Termín
- upravené výkresy pošleme do 26. 3. 2026 (dle zápisu z KD č. 3)

Cena víceprací (dle sazby ze smlouvy {sazba} Kč/h bez DPH)
- detailing {hod[0]} h, statika {hod[1]} h, celkem {sum(hod)} h
- celkem {cena_txt} Kč bez DPH

Pokud s rozsahem a cenou souhlasíte, prosím o potvrzení e-mailem, úpravy hned zahájíme.

S pozdravem
Ing. Karel Svoboda
STATIKA & DETAIL s.r.o. (fiktivní)

[Koncept od Clauda – před odesláním zkontrolujte čísla, termíny a tón.]
""", encoding="utf-8")

        # ---------------- check against the answer key
        key = {r["puvodni_cesta"].removeprefix("zakazka_neporadek/"): r
               for r in csv.DictReader((OUT / "klic_zarazeni.csv").open(encoding="utf-8-sig"), delimiter=";")}
        chyby = []
        for f in inv:
            k = key[f["rel"]]
            got = (f["zakazka"], f["stav"]) + ((f["typ"],) if f["stav"] not in ("smeti", "nezarazeno") else ("",))
            exp = (k["zakazka"], k["stav"], k["typ"])
            if f["typ"] == "VYK":
                got += (f["cislo"], f["revize"]); exp += (k["cislo"], k["revize"])
            if got != exp:
                chyby.append(f"{f['rel']}: {got} != {exp}")
        assert len(key) == len(inv)
        print("\n".join(souhrn[:5]))
        print(f"Výkresů v protokolu: {len(vyk)}, listů: {sum(f['meta']['stran'] for f in vyk)}; úkolů z KD: {len(ukoly)};"
              f" cena víceprací {cena} Kč")
        print("Varování:", *varovani, sep="\n  ")
        if chyby:
            raise SystemExit("NESOUHLASÍ S KLÍČEM:\n" + "\n".join(chyby))
        print("Zařazení souhlasí s klíčem (klic_zarazeni.csv).")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _fix_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            d = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                d = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-03-20T08:00:00Z", d)
            zi = zipfile.ZipInfo(info.filename, (2026, 3, 20, 8, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, d)
    return out.getvalue()


if __name__ == "__main__":
    main()
