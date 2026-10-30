"""Vzorové řešení cvičení „Data z Tekly“ (cviceni.qmd).

Čte syntetické exporty z data/tekla/ (vytvoří je tools/make_tekla_data.py),
originály nemění a výsledky ukládá do demos/tekla_data/.

Spuštění (z kořene workshopu):
    uv run --with pandas --with openpyxl --with matplotlib python tools/solve_tekla_data.py
"""
from __future__ import annotations

import base64
import html
import io
import math
import re
from pathlib import Path

import matplotlib
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "tekla"
OUT = ROOT / "demos" / "tekla_data"
OUT.mkdir(parents=True, exist_ok=True)

BLUE, INK, MUTED = "#2a78d6", "#1f2a37", "#52514e"
plt.rcParams.update({"axes.spines.top": False, "axes.spines.right": False, "font.size": 11,
                     "axes.edgecolor": "#b9b8b2", "axes.labelcolor": INK, "xtick.color": MUTED,
                     "ytick.color": MUTED, "axes.titlecolor": INK})
FILL = {"pridano": "C6EFCE", "odebrano": "FFC7CE", "zmeneno": "FFEB9C", "kontrola": "F4B183"}
STOCK, KERF = 12000, 3  # mm: obchodní délka tyče a šířka řezu


# ------------------------------------------------------------------ načtení
def read_tekla(path: Path) -> pd.DataFrame:
    """Najde řádek s hlavičkou tabulky, vynechá hlavičku reportu i řádek CELKEM."""
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("PHASE;"))
    body = [lines[start]]
    for l in lines[start + 1:]:
        if not l.strip() or l.startswith("CELKEM"):
            break
        body.append(l)
    df = pd.read_csv(io.StringIO("\n".join(body)), sep=";", decimal=",", dtype=str, keep_default_na=False)
    for c in ("PHASE", "QTY", "LENGTH", "CLASS"):
        if c in df:
            df[c] = df[c].astype(int)
    for c in ("WEIGHT", "WEIGHT_TOTAL", "AREA"):
        if c in df:
            df[c] = df[c].str.replace(",", ".").astype(float)
    return df


def report_total(path: Path) -> float:
    last = [l for l in path.read_text(encoding="utf-8-sig").splitlines() if l.startswith("CELKEM")][0]
    return float([x for x in last.split(";") if "," in x][0].replace(",", "."))


parts = {r: read_tekla(DATA / f"hala_rev{r}_part_list.csv") for r in "AB"}
asms = {r: read_tekla(DATA / f"hala_rev{r}_assembly_list.csv") for r in "AB"}
for r in "AB":
    parts[r]["WEIGHT_SUM"] = parts[r].QTY * parts[r].WEIGHT
    parts[r]["AREA_SUM"] = parts[r].QTY * parts[r].AREA
pB, aB = parts["B"], asms["B"]
PHASES = {1: "1 hlavní rámy", 2: "2 střecha", 3: "3 stěny"}


def ptype(profile: str) -> str:
    m = re.match(r"[A-Z]+", profile)
    return m.group(0) if m else "?"


def grade(material: str) -> str:
    return material[:4] if material else ""


def autosize(ws, maxw=45):
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(maxw, max(9, w + 2))
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def picture(fig, name):
    fig.tight_layout()
    fig.savefig(OUT / name, dpi=150)
    b = io.BytesIO()
    fig.savefig(b, format="png", dpi=150)
    plt.close(fig)
    return f'<img alt="{name}" src="data:image/png;base64,{base64.b64encode(b.getvalue()).decode()}">'


def t(x, nd=0):
    """Číslo po česku: mezera jako oddělovač tisíců, desetinná čárka."""
    return f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")


# ------------------------------------------------------------------ 1 · přehled
overview = {
    "řádků v seznamu dílů": len(pB),
    "řádků v seznamu sestav": len(aB),
    "kusů dílů (součet QTY)": int(pB.QTY.sum()),
    "kusů sestav (součet QTY)": int(aB.QTY.sum()),
    "hmotnost z dílů [kg]": round(pB.WEIGHT_SUM.sum(), 1),
    "hmotnost ze sestav [kg]": round(aB.WEIGHT_TOTAL.sum(), 1),
    "řádek CELKEM v reportu [kg]": report_total(DATA / "hala_revB_part_list.csv"),
}

# ------------------------------------------------------------------ 2 · souhrn hmotností
pB["PROFILE_TYPE"] = pB.PROFILE.map(ptype)
pB["GRADE"] = pB.MATERIAL.map(grade).replace("", "(chybí)")
by_phase = pB.groupby("PHASE").agg(kusu=("QTY", "sum"), hmotnost_kg=("WEIGHT_SUM", "sum")).reset_index()
by_phase["PHASE"] = by_phase.PHASE.map(PHASES)
by_profile = (pB.groupby(["PROFILE_TYPE", "PROFILE"]).agg(kusu=("QTY", "sum"), hmotnost_kg=("WEIGHT_SUM", "sum"))
              .reset_index().sort_values("hmotnost_kg", ascending=False))
by_type = pB.groupby("PROFILE_TYPE").WEIGHT_SUM.sum().sort_values()
by_mat = pB.groupby("MATERIAL").agg(kusu=("QTY", "sum"), hmotnost_kg=("WEIGHT_SUM", "sum")).reset_index()
by_mat["MATERIAL"] = by_mat.MATERIAL.replace("", "(chybí)")
total = pB.WEIGHT_SUM.sum()
for df in (by_phase, by_profile, by_mat):
    df["podil_%"] = (100 * df.hmotnost_kg / total).round(1)
    df["hmotnost_kg"] = df.hmotnost_kg.round(1)

fig, ax = plt.subplots(1, 2, figsize=(11, 4.2), gridspec_kw={"width_ratios": [1, 1.4]})
bars = ax[0].bar(by_phase.PHASE, by_phase.hmotnost_kg / 1000, color=BLUE, width=0.55)
ax[0].bar_label(bars, labels=[t(v / 1000, 1) + " t" for v in by_phase.hmotnost_kg], padding=3, color=INK)
ax[0].set(ylabel="Hmotnost (t)", title="Hmotnost podle fáze")
ax[0].grid(axis="y", color="#e6e5e0")
ax[0].set_axisbelow(True)
bars = ax[1].barh(by_type.index, by_type.values / 1000, color=BLUE, height=0.6)
ax[1].bar_label(bars, labels=[t(v / 1000, 1) + " t" for v in by_type.values], padding=3, color=INK)
ax[1].set(xlabel="Hmotnost (t)", title="Hmotnost podle typu profilu")
ax[1].grid(axis="x", color="#e6e5e0")
ax[1].set_axisbelow(True)
ax[1].set_xlim(0, by_type.max() / 1000 * 1.18)
img_weights = picture(fig, "hmotnosti.png")

with pd.ExcelWriter(OUT / "souhrn_hmotnosti.xlsx", engine="openpyxl") as w:
    pd.DataFrame(list(overview.items()), columns=["položka", "hodnota"]).to_excel(w, sheet_name="Přehled", index=False)
    by_phase.to_excel(w, sheet_name="Podle fáze", index=False)
    by_profile.to_excel(w, sheet_name="Podle profilu", index=False)
    by_mat.to_excel(w, sheet_name="Podle materiálu", index=False)
    pd.pivot_table(pB, index="PROFILE", columns="PHASE", values="WEIGHT_SUM", aggfunc="sum",
                   fill_value=0, margins=True, margins_name="Celkem").round(1).to_excel(w, sheet_name="Profil × fáze")
wb = load_workbook(OUT / "souhrn_hmotnosti.xlsx")
for ws in wb.worksheets:
    autosize(ws)
wb.save(OUT / "souhrn_hmotnosti.xlsx")

# ------------------------------------------------------------------ 3 · kontrola reportu
problems = []


def flag(kind, df, why, cols=("ASSEMBLY_POS", "PART_POS", "PROFILE", "MATERIAL", "LENGTH", "QTY")):
    for _, r in df.iterrows():
        problems.append({"problém": kind, **{c: r.get(c, "") for c in cols}, "vysvětlení": why})


valid_profile = re.compile(r"^(HEA|HEB|IPE|UPE|UPN)\d+$|^(RHS|SHS)\d+\*\d+\*\d+$|^L\d+\*\d+$|^D\d+$|^PL\d+\*\d+$")
flag("Chybí materiál", pB[pB.MATERIAL == ""],
     "Bez jakosti nelze správně nacenit ani objednat materiál; v modelu doplnit materiál dílu.")
flag("Neočíslováno (?)", pB[(pB.ASSEMBLY_POS.str.contains(r"\?")) | (pB.PART_POS.str.contains(r"\?"))],
     "Číslování není aktuální – před exportem spustit Numbering, jinak díl nemá výkres ani značku.")
flag("Prázdné UDA_VYKRES", pB[pB.UDA_VYKRES == ""], "Sestava nemá přiřazený výrobní výkres.")
flag("Prázdné UDA_DODAVKA", pB[pB.UDA_DODAVKA == ""], "Není jasné, kdo díl dodává (OCEL / jiný dodavatel).")
flag("Neznámý profil (překlep?)", pB[~pB.PROFILE.str.match(valid_profile)],
     "Profil neodpovídá katalogovému tvaru – pravděpodobně HEA200; opravit v modelu.")
num = aB[aB.ASSEMBLY_POS != "?"]
dups = num[num.duplicated("ASSEMBLY_POS", keep=False)]
dups = dups[dups.groupby("ASSEMBLY_POS").LENGTH.transform("nunique") > 1]
flag("Duplicitní značka sestavy", dups,
     "Stejná značka pro sestavy s různou délkou – na dílně by se zaměnily; přečíslovat.",
     cols=("ASSEMBLY_POS", "PROFILE", "MATERIAL", "LENGTH", "QTY"))
pdups = pB[pB.PART_POS != "?"].groupby("PART_POS").filter(lambda g: g[["PROFILE", "LENGTH"]].drop_duplicates().shape[0] > 1)
flag("Duplicitní značka dílu", pdups, "Jedna pozice dílu se dvěma různými délkami.")
# Konzistence: hmotnost sestavy = součet dílů
chk = pB.groupby("ASSEMBLY_POS").WEIGHT_SUM.sum().round(1)
chk2 = aB.groupby("ASSEMBLY_POS").WEIGHT_TOTAL.sum().round(1)
consistent = bool((chk - chk2).abs().max() < 0.5)
problems_df = pd.DataFrame(problems).fillna("")
with pd.ExcelWriter(OUT / "kontrola_reportu.xlsx", engine="openpyxl") as w:
    problems_df.to_excel(w, sheet_name="Problémy", index=False)
    problems_df.groupby("problém").size().rename("řádků").reset_index().to_excel(w, sheet_name="Souhrn", index=False)
wb = load_workbook(OUT / "kontrola_reportu.xlsx")
for ws in wb.worksheets:
    autosize(ws, 70)
wb.save(OUT / "kontrola_reportu.xlsx")
problem_summary = problems_df.groupby("problém", sort=False).size()

# ------------------------------------------------------------------ 4 · porovnání revizí
FIELDS = ["QTY", "PROFILE", "MATERIAL", "LENGTH", "WEIGHT"]


def compare(a: pd.DataFrame, b: pd.DataFrame, key: list[str]):
    rows = []
    keys = list(dict.fromkeys(list(a[key].itertuples(index=False, name=None))
                              + list(b[key].itertuples(index=False, name=None))))
    for k in keys:
        ra = a.loc[(a[key] == pd.Series(k, index=key)).all(axis=1)]
        rb = b.loc[(b[key] == pd.Series(k, index=key)).all(axis=1)]
        base = dict(zip(key, k))
        if any("?" in str(x) for x in k):
            status, detail = "kontrola", "neočíslováno (?) – nelze spárovat, nejspíš nová položka"
        elif ra.empty:
            status, detail = "pridano", "nová položka"
        elif rb.empty:
            status, detail = "odebrano", "položka zrušena"
        elif len(ra) > 1 or len(rb) > 1:
            status, detail = "kontrola", f"duplicitní značka (řádků: A {len(ra)}, B {len(rb)})"
        else:
            diffs = [f"{f}: {ra.iloc[0][f]} → {rb.iloc[0][f]}" for f in FIELDS if f in a and ra.iloc[0][f] != rb.iloc[0][f]]
            if not diffs:
                continue
            status, detail = "zmeneno", "; ".join(diffs)
        wa = (ra.QTY * ra.WEIGHT).sum()
        wb_ = (rb.QTY * rb.WEIGHT).sum()
        rows.append({**base, "stav": status, "co se změnilo": detail,
                     "hmotnost A [kg]": round(wa, 1), "hmotnost B [kg]": round(wb_, 1),
                     "rozdíl [kg]": round(wb_ - wa, 1)})
    return pd.DataFrame(rows)


diff_asm = compare(asms["A"], asms["B"], ["ASSEMBLY_POS"])
diff_part = compare(parts["A"], parts["B"], ["ASSEMBLY_POS", "PART_POS"])
STATUS_CZ = {"pridano": "přidáno", "odebrano": "odebráno", "zmeneno": "změněno", "kontrola": "zkontrolovat"}
wA, wB = parts["A"].WEIGHT_SUM.sum(), parts["B"].WEIGHT_SUM.sum()
with pd.ExcelWriter(OUT / "porovnani_revA_revB.xlsx", engine="openpyxl") as w:
    pd.DataFrame([["Revize A (12. 10. 2026)", round(wA, 1)], ["Revize B (27. 10. 2026)", round(wB, 1)],
                  ["Rozdíl B − A", round(wB - wA, 1)]], columns=["", "hmotnost [kg]"]).to_excel(w, sheet_name="Souhrn", index=False)
    diff_asm.to_excel(w, sheet_name="Sestavy", index=False)
    diff_part.to_excel(w, sheet_name="Díly", index=False)
wb = load_workbook(OUT / "porovnani_revA_revB.xlsx")
for name in ("Sestavy", "Díly"):
    ws = wb[name]
    scol = [c.value for c in ws[1]].index("stav") + 1
    for row in ws.iter_rows(min_row=2):
        st = row[scol - 1].value
        for c in row:
            c.fill = PatternFill("solid", fgColor=FILL[st])
        row[scol - 1].value = STATUS_CZ[st]
for ws in wb.worksheets:
    autosize(ws, 80)
wb["Souhrn"]["A6"] = "Barvy: zelená = přidáno, červená = odebráno, žlutá = změněno, oranžová = zkontrolovat ručně."
wb.save(OUT / "porovnani_revA_revB.xlsx")

# ------------------------------------------------------------------ 5 · cenová nabídka
cen = pd.ExcelFile(DATA / "cenik.xlsx")
mat_price = pd.read_excel(cen, "Material", header=2).dropna(subset=["S235 [Kc/kg]"])
mat_price = mat_price.set_index("Typ profilu")
finish = pd.read_excel(cen, "Povrch", header=2).set_index("FINISH")
other = pd.read_excel(cen, "Prace a ostatni").set_index("Polozka")["Hodnota"]
TYPE_FIX = {"HAE": "HEA", "SHS": "RHS", "HEB": "HEA"}
assumptions = ["Profil HAE200 (překlep v reportu) naceněn jako HEA200.",
               "Díly bez materiálu (výztuhy 1104, plechy překladu „?“) naceněny jako S235.",
               "Překlad vrat s pozicí „?“ je v nabídce zahrnut (je v modelu, jen neočíslovaný)."]
pr = pB.copy()
pr["TYP"] = pr.PROFILE_TYPE.replace(TYPE_FIX)
pr["JAKOST"] = pr.MATERIAL.map(grade).replace("", "S235")
pr["KC_KG"] = [mat_price.loc[ty, f"{g} [Kc/kg]"] for ty, g in zip(pr.TYP, pr.JAKOST)]
pr["MATERIAL_KC"] = pr.WEIGHT_SUM * pr.KC_KG
mat_tab = pr.groupby(["TYP", "JAKOST"]).agg(hmotnost_kg=("WEIGHT_SUM", "sum"), cena_kc_kg=("KC_KG", "first"),
                                            cena_kc=("MATERIAL_KC", "sum")).reset_index()
W = pr.WEIGHT_SUM.sum()
mat_cost = pr.MATERIAL_KC.sum()
prorez = mat_cost * other["Prorez materialu"] / 100
spoj = (mat_cost + prorez) * other["Spojovaci material"] / 100
paint_area = pr.loc[pr.FINISH == "NATER_C3", "AREA_SUM"].sum()
zinc_w = pr.loc[pr.FINISH == "ZINEK", "WEIGHT_SUM"].sum()
trips = math.ceil(W / 20000)
lines = [
    ("Materiál – ocel podle ceníku", f"{t(W)} kg", mat_cost),
    (f"Prořez materiálu {other['Prorez materialu']:.0f} %", "", prorez),
    (f"Spojovací materiál {other['Spojovaci material']:.0f} %", "", spoj),
    ("Dílenská výroba", f"{t(W)} kg × {other['Dilenska vyroba']:.0f} Kč/kg", W * other["Dilenska vyroba"]),
    ("Nátěr C3 (NATER_C3)", f"{t(paint_area, 1)} m² × {finish.loc['NATER_C3', 'Cena']:.0f} Kč/m²",
     paint_area * finish.loc["NATER_C3", "Cena"]),
    ("Žárové zinkování (ZINEK)", f"{t(zinc_w)} kg × {finish.loc['ZINEK', 'Cena']:.0f} Kč/kg",
     zinc_w * finish.loc["ZINEK", "Cena"]),
    ("Montáž", f"{t(W)} kg × {other['Montaz']:.0f} Kč/kg", W * other["Montaz"]),
    ("Doprava", f"{trips} jízdy × {t(other['Doprava'])} Kč", trips * other["Doprava"]),
]
offer = pd.DataFrame(lines, columns=["položka", "výpočet", "cena bez DPH [Kč]"])
offer["cena bez DPH [Kč]"] = offer["cena bez DPH [Kč]"].round(0)
net = offer["cena bez DPH [Kč]"].sum()
vat = round(net * other["DPH"] / 100)
gross = net + vat
with pd.ExcelWriter(OUT / "nabidka.xlsx", engine="openpyxl") as w:
    offer.to_excel(w, sheet_name="Nabídka", index=False, startrow=3)
    mat_tab.round(1).to_excel(w, sheet_name="Materiál podle typu", index=False)
    pd.DataFrame({"předpoklad": assumptions}).to_excel(w, sheet_name="Předpoklady", index=False)
wb = load_workbook(OUT / "nabidka.xlsx")
ws = wb["Nabídka"]
ws["A1"] = "Cenová nabídka – ocelová konstrukce HALA SKLAD 24x48 (zakázka Z-2026-014, revize B)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Návrh pro interní kontrolu · syntetická data · ceny bez záruky"
n = ws.max_row
for label, val in [("Celkem bez DPH", net), (f"DPH {other['DPH']:.0f} %", vat), ("Celkem s DPH", gross)]:
    n += 1
    ws.cell(n, 1, label).font = Font(bold=True)
    ws.cell(n, 3, val).font = Font(bold=True)
for row in ws.iter_rows(min_row=5, min_col=3, max_col=3):
    row[0].number_format = '# ##0" Kč"'
for sh in wb.worksheets:
    autosize(sh, 60)
    sh.freeze_panes = None
wb.save(OUT / "nabidka.xlsx")

offer_rows = "".join(f"<tr><td>{html.escape(a)}</td><td>{html.escape(b)}</td><td class=n>{t(c)} Kč</td></tr>"
                     for a, b, c in offer.itertuples(index=False))
mat_rows = "".join(f"<tr><td>{r.TYP}</td><td>{r.JAKOST}</td><td class=n>{t(r.hmotnost_kg)}</td>"
                   f"<td class=n>{t(r.cena_kc_kg, 1)}</td><td class=n>{t(r.cena_kc)} Kč</td></tr>"
                   for r in mat_tab.itertuples())
(OUT / "nabidka.html").write_text(f"""<!doctype html><html lang="cs"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Cenová nabídka – hala</title>
<style>
body{{font:15px/1.5 system-ui,sans-serif;max-width:820px;margin:auto;padding:32px 20px;color:#1f2a37;background:#fff}}
header{{display:flex;justify-content:space-between;gap:16px;border-bottom:3px solid #12355b;padding-bottom:12px}}
h1{{margin:0;font-size:26px;color:#12355b}} .muted{{color:#5b6573}} table{{width:100%;border-collapse:collapse;margin:14px 0}}
td,th{{padding:7px 8px;border-bottom:1px solid #e3e6ea;text-align:left;vertical-align:top}} th{{background:#f1f5f9}}
.n{{text-align:right;white-space:nowrap}} tfoot td{{font-weight:700;border-top:2px solid #12355b}}
.big td{{font-size:18px}} .note{{background:#fff7e0;border-left:4px solid #e0a100;padding:10px 14px;margin:18px 0}}
@media print{{body{{padding:0}} .note{{break-inside:avoid}}}}
</style>
<header><div><h1>Cenová nabídka</h1><div class=muted>Ocelová konstrukce · HALA SKLAD 24×48 m</div></div>
<div class="muted n">Zakázka Z-2026-014<br>Podklad: model revize B (27. 10. 2026)<br>Návrh – NEODESÍLAT bez kontroly</div></header>
<p>Rozsah: sloupy, příčle, vaznice, paždíky, ztužidla a rám vrat podle exportu z Tekla Structures,
celková hmotnost <b>{t(W)} kg</b>. Ceny podle interního ceníku (syntetický), bez DPH.</p>
<table><thead><tr><th>Položka</th><th>Výpočet</th><th class=n>Cena bez DPH</th></tr></thead><tbody>{offer_rows}</tbody>
<tfoot><tr><td colspan=2>Celkem bez DPH</td><td class=n>{t(net)} Kč</td></tr>
<tr><td colspan=2>DPH {other['DPH']:.0f} %</td><td class=n>{t(vat)} Kč</td></tr>
<tr class=big><td colspan=2>Celkem s DPH</td><td class=n>{t(gross)} Kč</td></tr></tfoot></table>
<h2>Materiál podle typu profilu</h2>
<table><thead><tr><th>Typ</th><th>Jakost</th><th class=n>Hmotnost [kg]</th><th class=n>Kč/kg</th><th class=n>Cena</th></tr></thead>
<tbody>{mat_rows}</tbody></table>
<div class=note><b>Předpoklady a otevřené body</b><ul>{''.join(f'<li>{html.escape(a)}</li>' for a in assumptions)}
<li>Nezahrnuje: kotvy a chemické kotvení, opláštění, projekt, statiku, jeřáb nad rámec montáže.</li></ul></div>
<p class=muted>Výuková ukázka – data i ceny jsou vymyšlené. Vytvořeno skriptem tools/solve_tekla_data.py.</p>
</html>""", encoding="utf-8")

# ------------------------------------------------------------------ 6 · bonus: nářezový plán a NC
bars_src = pB[~pB.PROFILE.str.startswith("PL")].copy()
bars_src["PROFILE_OK"] = bars_src.PROFILE.replace({"HAE200": "HEA200"})
cut_summary, cut_patterns = [], []
for prof, g in bars_src.groupby("PROFILE_OK"):
    pieces = sorted([L for L, q in zip(g.LENGTH, g.QTY) for _ in range(q)], reverse=True)
    bins: list[list[int]] = []
    for L in pieces:  # first-fit decreasing
        for b in bins:
            if sum(b) + KERF * len(b) + L <= STOCK:
                b.append(L)
                break
        else:
            bins.append([L])
    used = sum(pieces)
    cut_summary.append({"profil": prof, "kusů": len(pieces), "součet délek [m]": round(used / 1000, 2),
                        "tyčí 12 m": len(bins), "využití [%]": round(100 * used / (len(bins) * STOCK), 1),
                        "odpad [m]": round((len(bins) * STOCK - used) / 1000, 2)})
    pats = pd.Series([" + ".join(map(str, b)) for b in bins]).value_counts()
    for pat, cnt in pats.items():
        rest = STOCK - sum(map(int, pat.split(" + "))) - KERF * (len(pat.split(" + ")) - 1)
        cut_patterns.append({"profil": prof, "počet tyčí": cnt, "řezy [mm]": pat, "zbytek [mm]": rest})
cut_summary = pd.DataFrame(cut_summary)
with pd.ExcelWriter(OUT / "narezovy_plan.xlsx", engine="openpyxl") as w:
    cut_summary.to_excel(w, sheet_name="Souhrn", index=False)
    pd.DataFrame(cut_patterns).to_excel(w, sheet_name="Plán řezů", index=False)
wb = load_workbook(OUT / "narezovy_plan.xlsx")
for ws in wb.worksheets:
    autosize(ws, 60)
wb.save(OUT / "narezovy_plan.xlsx")


def read_nc1(path: Path) -> dict:
    lines = [l.strip() for l in path.read_text(encoding="ascii").splitlines()]
    head = [l for l in lines[1:lines.index("EN")] if not l.startswith("**")]
    head = head[:head.index("BO")] if "BO" in head else head
    holes = sum(1 for l in lines if l.startswith("v "))
    return dict(PART_POS=head[3], MATERIAL=head[4], QTY=int(head[5]), PROFILE=head[6],
                LENGTH=float(head[8]), holes=holes, soubor=path.name)


nc = pd.DataFrame([read_nc1(p) for p in sorted((DATA / "nc_dstv").glob("*.nc1"))])
need = pB[~pB.PROFILE.str.startswith("PL") & (pB.PROFILE != "D20")]
need_g = need.groupby("PART_POS").agg(PROFILE=("PROFILE", lambda s: ", ".join(sorted(set(s)))),
                                      LENGTH=("LENGTH", lambda s: ", ".join(map(str, sorted(set(s))))),
                                      QTY=("QTY", "sum")).reset_index()
nc_rows = []
for _, r in need_g.iterrows():
    m = nc[nc.PART_POS == r.PART_POS]
    if r.PART_POS == "?":
        nc_rows.append({**r, "stav": "neočíslovaný díl – NC nelze vytvořit, nejdřív očíslovat"})
        continue
    if m.empty:
        nc_rows.append({**r, "stav": "CHYBÍ NC soubor"})
        continue
    m = m.iloc[0]
    issues = []
    if m.PROFILE != r.PROFILE:
        issues.append(f"profil v NC {m.PROFILE} ≠ report {r.PROFILE}")
    if str(int(m.LENGTH)) != r.LENGTH:
        issues.append(f"délka v NC {int(m.LENGTH)} ≠ report {r.LENGTH}")
    if m.QTY != r.QTY:
        issues.append(f"kusů v NC {m.QTY} ≠ report {r.QTY}")
    nc_rows.append({**r, "stav": "; ".join(issues) or "OK"})
for _, m in nc[~nc.PART_POS.isin(need_g.PART_POS)].iterrows():
    nc_rows.append({"PART_POS": m.PART_POS, "PROFILE": m.PROFILE, "LENGTH": str(int(m.LENGTH)), "QTY": m.QTY,
                    "stav": "NC soubor navíc – díl už v reportu není (starý export?)"})
nc_check = pd.DataFrame(nc_rows)
with pd.ExcelWriter(OUT / "kontrola_nc.xlsx", engine="openpyxl") as w:
    nc_check.to_excel(w, sheet_name="Kontrola NC", index=False)
wb = load_workbook(OUT / "kontrola_nc.xlsx")
ws = wb.active
for row in ws.iter_rows(min_row=2):
    if row[-1].value != "OK":
        for c in row:
            c.fill = PatternFill("solid", fgColor=FILL["zmeneno"])
autosize(ws, 70)
wb.save(OUT / "kontrola_nc.xlsx")

# ------------------------------------------------------------------ souhrnný HTML report
def table(df, fmt=None):
    return df.to_html(index=False, border=0, float_format=lambda x: t(x, 1), na_rep="")


diff_view = diff_asm.copy()
diff_view["stav"] = diff_view.stav.map(STATUS_CZ)
(OUT / "souhrn.html").write_text(f"""<!doctype html><html lang="cs"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Souhrn z reportu Tekla</title>
<style>body{{font:16px/1.55 system-ui,sans-serif;max-width:1000px;margin:auto;padding:24px;color:#1f2a37;background:#fff}}
img{{max-width:100%}} table{{border-collapse:collapse;display:block;overflow:auto;margin:10px 0 18px}}
td,th{{padding:6px 9px;border-bottom:1px solid #e3e6ea;text-align:left}} th{{background:#f1f5f9}}
h1,h2{{line-height:1.2;color:#12355b}} .note{{background:#edf6fc;padding:14px;border-radius:10px}}
.kpi{{display:flex;flex-wrap:wrap;gap:12px}} .kpi div{{background:#f1f5f9;border-radius:10px;padding:10px 16px}}
.kpi b{{display:block;font-size:24px;color:#12355b}}</style>
<h1>Souhrn z reportu Tekla – hala 24 × 48 m, revize B</h1>
<p class=note>Výuková ukázka · všechna data jsou syntetická. Vytvořeno skriptem <code>tools/solve_tekla_data.py</code>.</p>
<div class=kpi><div><b>{t(total, 1)} kg</b>celková hmotnost</div><div><b>{overview['kusů dílů (součet QTY)']}</b>kusů dílů</div>
<div><b>{overview['kusů sestav (součet QTY)']}</b>kusů sestav</div><div><b>{len(problems_df)}</b>řádků s problémem</div>
<div><b>{t(wB - wA, 1)} kg</b>rozdíl proti revizi A</div></div>
<h2>1 · Přehled exportu</h2>{table(pd.DataFrame(list(overview.items()), columns=['položka', 'hodnota']))}
<p>Hmotnost sestav souhlasí se součtem dílů: <b>{'ano' if consistent else 'NE'}</b>. Hmotnost v reportu je za 1 kus – vždy násobit QTY.</p>
<h2>2 · Hmotnosti</h2>{img_weights}{table(by_phase)}{table(by_mat)}
<h2>3 · Kontrola reportu</h2>{table(problem_summary.rename('řádků').reset_index())}{table(problems_df)}
<h2>4 · Revize A → B (sestavy)</h2>{table(diff_view)}
<p>Revize A {t(wA, 1)} kg → revize B {t(wB, 1)} kg.</p>
<h2>5 · Cenová nabídka</h2><p>Celkem bez DPH <b>{t(net)} Kč</b>, s DPH {t(gross)} Kč. Podrobně v <a href="nabidka.html">nabidka.html</a>.</p>
<h2>6 · Nářezový plán (tyče 12 m, řez 3 mm, first-fit decreasing)</h2>{table(cut_summary)}
<h2>6b · Kontrola NC souborů</h2>{table(nc_check)}
</html>""", encoding="utf-8")

# ------------------------------------------------------------------ klíčová čísla pro cviceni.qmd
print("PREHLED", overview)
print("FAZE\n", by_phase.to_string(index=False))
print("MATERIAL\n", by_mat.to_string(index=False))
print("PROFIL\n", by_profile.to_string(index=False))
print("PROBLEMY\n", problem_summary.to_string())
print("DIFF SESTAVY\n", diff_asm.to_string(index=False))
print(f"A {wA:.1f}  B {wB:.1f}  rozdil {wB - wA:.1f}")
print("NABIDKA\n", offer.to_string(index=False), f"\nnetto {net}  DPH {vat}  brutto {gross}")
print(f"natirana plocha {paint_area:.1f} m2, zinek {zinc_w:.1f} kg, jizdy {trips}")
print("REZ\n", cut_summary.to_string(index=False), "\ncelkem tyci", cut_summary["tyčí 12 m"].sum())
print("NC\n", nc_check.to_string(index=False))
