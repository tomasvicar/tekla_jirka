"""Ukázkové řešení: roztřídění výkresů vytištěných z Tekly do složek zakázky.

Vstup:  data/pdf_export/ (vytvoří tools/make_pdf_export_data.py)
Výstup: demos/ukazky/pdf-export/strom.txt              strom vytvořených složek
        demos/ukazky/pdf-export/predavaci_protokol.xlsx předávací protokol + list Problémy
        (roztříděné PDF se skládají do dočasné složky, na web nejdou)

Struktura:
    Z2026-021_Hala_Lipnik/04_Vykresy/
        00_Komplet/        jedno PDF za každý druh výkresu (spojené)
        01_Vykresy_rozmisteni/   G01_KOTEVNI_PLAN_revB.pdf ...
        02_Sestavy/              A03_PRICLE_revA.pdf ...
        03_Dilce/                P02_CELNI_DESKA_revA.pdf ...
        _Archiv/           starší revize

Spuštění (z kořene workshopu):
    uv run --with pypdf --with openpyxl python tools/solve_pdf_export.py
"""
from __future__ import annotations

import csv
import hashlib
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "pdf_export"
DEMO = ROOT / "demos" / "ukazky" / "pdf-export"
ORDER, PROJECT, HANDOVER = "Z2026-021", "Hala sklad Lipník", "2026-10-30"

FOLDERS = {"G": "01_Vykresy_rozmisteni", "A": "02_Sestavy", "W": "03_Dilce"}
KIND = {"G": "výkres rozmístění", "A": "sestava", "W": "dílec"}
# „A3 - PRICLE - A.pdf“ -> značka A3, název PRICLE, revize A (revize nemusí být)
NAME_RE = re.compile(r"^([AGP])(\d+) - ([^()]+?)(?: - ([A-Z]))?\.pdf$")


def target_name(mark: str, name: str, rev: str) -> str:
    letter, num = mark.split(".")
    return f"{letter}{int(num):02d}_{name.replace(' ', '_')}" + (f"_rev{rev}" if rev else "") + ".pdf"


def main() -> None:
    with open(SRC / "seznam_vykresu.csv", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    listed = {r["Znacka"]: r for r in rows}

    files, junk, dupes, problems = {}, [], [], []
    hashes: dict[str, str] = {}
    # Nejdřív soubory s rozpoznaným názvem, pak ostatní (kvůli hledání kopií).
    for p in sorted((SRC / "tisk_z_tekly").iterdir(), key=lambda p: (not NAME_RE.match(p.name), p.name)):
        h = hashlib.md5(p.read_bytes()).hexdigest()
        m = NAME_RE.match(p.name)
        if p.suffix.lower() != ".pdf":
            junk.append(p.name)
        elif not m:
            if h in hashes:
                dupes.append((p.name, hashes[h]))
            else:
                problems.append((p.name, "Nerozpoznaný název souboru"))
        else:
            mark = f"{m[1]}.{m[2]}"
            files[(mark, m[4] or "")] = p
            hashes[h] = p.name
            # Kontrola obsahu: sedí značka v razítku s názvem souboru?
            text = PdfReader(p).pages[0].extract_text()
            if mark not in text:
                problems.append((p.name, f"V razítku není značka {mark}"))
    for name, orig in dupes:
        problems.append((name, f"Kopie souboru {orig} – nepoužito"))
    for name in junk:
        problems.append((name, "Není PDF – nepoužito"))

    work = Path(tempfile.mkdtemp())
    out = work / "Z2026-021_Hala_Lipnik" / "04_Vykresy"
    for d in [*FOLDERS.values(), "_Archiv", "00_Komplet"]:
        (out / d).mkdir(parents=True)

    protocol, archived = [], 0
    for (mark, rev), p in sorted(files.items()):
        row = listed.get(mark)
        if row is None:
            problems.append((p.name, "Výkres není v seznamu výkresů"))
            continue
        current = rev == row["Revize"]
        dest = out / (FOLDERS[row["Typ"]] if current else "_Archiv") / target_name(mark, row["Nazev"], rev)
        shutil.copy2(p, dest)
        if current:
            note = "NEAKTUÁLNÍ v Tekle – přetisknout" if row["Stav"] != "Aktuální" else ""
            protocol.append((row, dest, note))
        else:
            archived += 1
    printed = {m for (m, _r) in files}
    for mark, row in listed.items():
        if (mark, row["Revize"]) not in files:
            what = f"revize {row['Revize']}" if row["Revize"] else "výkres"
            problems.append((mark, f"Chybí PDF ({what}) – vytisknout z Tekly"
                             + (" (ve složce je jen starší revize)" if mark in printed else "")))
        elif row["Stav"] != "Aktuální":
            problems.append((mark, "Výkres je v Tekle neaktuální – aktualizovat a přetisknout"))

    # Spojené PDF za každý druh výkresu.
    for typ, folder in FOLDERS.items():
        w = PdfWriter()
        for p in sorted((out / folder).glob("*.pdf")):
            w.append(str(p))
        with open(out / "00_Komplet" / f"{folder}_komplet.pdf", "wb") as f:
            w.write(f)

    # Strom složek jako text.
    lines = [out.parent.name + "/", "└── 04_Vykresy/"]
    dirs = sorted(d for d in out.iterdir())
    for i, d in enumerate(dirs):
        last = i == len(dirs) - 1
        items = sorted(x.name for x in d.iterdir())
        lines.append(f"    {'└──' if last else '├──'} {d.name}/  ({len(items)} PDF)")
        for j, name in enumerate(items):
            lines.append(f"    {'    ' if last else '│   '}{'└──' if j == len(items) - 1 else '├──'} {name}")
    DEMO.mkdir(parents=True, exist_ok=True)
    (DEMO / "strom.txt").write_text("\n".join(lines) + "\n", encoding="utf-8-sig")  # BOM: prohlížeč pozná češtinu

    write_protocol(protocol, problems)
    counts = {f: len(list((out / f).iterdir())) for f in [*FOLDERS.values(), "_Archiv"]}
    shutil.rmtree(work)
    print("Roztříděno:", counts)
    print(f"Protokol: {len(protocol)} výkresů, problémů: {len(problems)}")
    for name, why in problems:
        print(f"  {name}: {why}")


def write_protocol(protocol, problems) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Protokol"
    bold, thin = Font(bold=True), Side(style="thin", color="999999")
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    warn_fill = PatternFill("solid", fgColor="FCE4D6")
    ws["A1"] = "PŘEDÁVACÍ PROTOKOL – VÝKRESOVÁ DOKUMENTACE"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["C2"] = "Zakázka:", f"{ORDER} {PROJECT}"
    ws["A3"], ws["C3"] = "Datum předání:", HANDOVER
    ws["A4"], ws["C4"] = "Počet výkresů:", len(protocol)
    header = ["Č.", "Značka", "Název", "Druh", "Revize", "Datum revize", "Fáze", "Soubor", "Poznámka"]
    ws.append([])
    ws.append(header)
    for c in ws[6]:
        c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
    order = {"G": 0, "A": 1, "W": 2}
    protocol.sort(key=lambda t: (order[t[0]["Typ"]], int(t[0]["Znacka"].split(".")[1])))
    for i, (row, dest, note) in enumerate(protocol, 1):
        ws.append([i, row["Znacka"], row["Nazev"], KIND[row["Typ"]], row["Revize"] or "–",
                   row["Datum_revize"], int(row["Faze"]), f"{dest.parent.name}/{dest.name}", note])
        if note:
            for c in ws[ws.max_row]:
                c.fill = warn_fill
    r = ws.max_row + 3
    ws.cell(r, 1, "Předal:"), ws.cell(r, 5, "Převzal:")
    ws.cell(r + 1, 1, "Podpis a datum: ____________"), ws.cell(r + 1, 5, "Podpis a datum: ____________")
    for col, w in zip("ABCDEFGHI", [5, 9, 20, 18, 8, 13, 6, 46, 34]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "6:6"

    wp = wb.create_sheet("Problémy")
    wp.append(["Soubor / výkres", "Problém"])
    for c in wp[1]:
        c.font, c.fill = bold, head_fill
    for name, why in problems:
        wp.append([name, why])
    wp.column_dimensions["A"].width, wp.column_dimensions["B"].width = 28, 70
    for row in wp.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.properties.creator = "solve_pdf_export.py"
    wb.save(DEMO / "predavaci_protokol.xlsx")


if __name__ == "__main__":
    main()
